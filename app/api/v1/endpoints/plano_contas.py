# app/api/v1/endpoints/plano_contas.py

from collections import Counter
from typing import Any, Dict, List, Optional
import csv
import io
import re
import unicodedata

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from sqlmodel import Session, select, func
from loguru import logger
from pydantic import BaseModel
import openpyxl

from datetime import datetime
from app.db.session import get_db
from app.models.plano_contas import PlanoContas
from app.models.lancamento import Lancamento
from app.models.lancamento_cartao import LancamentoCartao
from app.models.integracao_bancaria import IntegracaoBancaria
from app.models.mapeamento_categoria import MapeamentoCategoria
from app.schemas.plano_contas import PlanoContasCreate, PlanoContasRead, PlanoContasUpdate
from app.crud import crud_plano_contas
from app.api.v1.deps import get_current_active_user, get_empresa_id_from_user, require_permission, require_any_permission
from app.models.usuario import Usuario

router = APIRouter()
MAX_PLANO_FILE_SIZE = 10 * 1024 * 1024


def _normalizar_tipo_plano(tipo: str) -> str:
    valor = (tipo or "").strip().upper()
    if valor.startswith("R"):
        return "R"
    if valor.startswith("D"):
        return "D"
    return "D"


def _normalizar_chave_nome(value: str) -> str:
    texto = unicodedata.normalize("NFKD", str(value or ""))
    texto = "".join(ch for ch in texto if not unicodedata.combining(ch))
    texto = re.sub(r"\s+", " ", texto).strip().upper()
    return texto


def _normalizar_codigo_hierarquia(value: Any) -> Optional[str]:
    raw = str(value or "").strip()
    if not raw:
        return None

    normalized = raw.replace("-", ".").replace("/", ".")
    parts = []
    for chunk in normalized.split("."):
        digits = re.sub(r"\D", "", chunk or "")
        if not digits:
            continue
        parts.append(digits.zfill(2))

    if not parts:
        return None
    return ".".join(parts)


def _parse_classificacao(raw_value: Any) -> tuple[Optional[str], str]:
    texto = str(raw_value or "").strip()
    if not texto:
        return None, ""

    match = re.match(r"^\s*([0-9][0-9\./-]*)\s*\.?\s*(.*)$", texto)
    if not match:
        return None, texto

    codigo = _normalizar_codigo_hierarquia(match.group(1))
    nome = (match.group(2) or "").strip(" .")
    if not nome:
        nome = texto
    return codigo, nome


def _normalizar_tipo_planilha(raw_value: Any) -> Optional[str]:
    texto = _normalizar_chave_nome(str(raw_value or ""))
    if not texto:
        return None
    if texto.startswith("R") or "RECEB" in texto or "ENTRADA" in texto or "RECEITA" in texto:
        return "R"
    if texto.startswith("P") or texto.startswith("D") or "PAG" in texto or "DESP" in texto or "SAIDA" in texto:
        return "D"
    return None


def _pick_header_index(headers: list[str], aliases: list[str]) -> Optional[int]:
    for alias in aliases:
        if alias in headers:
            return headers.index(alias)
    return None


def _is_true_like(raw_value: Any) -> Optional[bool]:
    text = _normalizar_chave_nome(str(raw_value or ""))
    if not text:
        return None
    if text in {"SIM", "TRUE", "1", "S", "YES"}:
        return True
    if text in {"NAO", "NAO ", "FALSE", "0", "N", "NO"}:
        return False
    return None

# --- SCHEMA LOCAL PARA REORDENAÇÃO ---
class ReordenacaoItem(BaseModel):
    id: int
    codigo: str
    conta_pai_id: Optional[int] = None
    tipo: str


class PlanoContasBulkCreateItem(BaseModel):
    nome: str
    tipo: str = "D"
    conta_pai_id: Optional[int] = None
    codigo: Optional[str] = None
    permite_lancamentos: Optional[bool] = True
    eh_operacional: Optional[bool] = True
    considerar_nos_resultados: Optional[bool] = True
    dre_grupo: Optional[str] = None


class PlanoContasBulkSyncItem(BaseModel):
    id: int
    nome: str
    tipo: str
    codigo: Optional[str] = None
    conta_pai_id: Optional[int] = None
    permite_lancamentos: Optional[bool] = True
    eh_operacional: Optional[bool] = True
    considerar_nos_resultados: Optional[bool] = True
    dre_grupo: Optional[str] = None


class PlanoContasBulkSyncPayload(BaseModel):
    items: List[PlanoContasBulkSyncItem]

# --- ENDPOINTS ---

@router.get(
    "/",
    response_model=List[PlanoContasRead],
    dependencies=[Depends(require_any_permission(["plano_contas:view", "page:lancamentos:view", "page:boletim:view", "page:dre:view"]))],
)
def read_plano_contas(
    *,
    db: Session = Depends(get_db),
    empresa_id: int = Depends(get_empresa_id_from_user),
):
    """Lista todas as categorias do plano de contas da empresa."""
    crud_plano_contas.normalize_company_operational_categories(db=db, empresa_id=empresa_id)
    # Ordena pelo código para garantir a árvore correta na leitura
    contas = db.exec(
        select(PlanoContas)
        .where(
            PlanoContas.empresa_id == empresa_id,
            PlanoContas.oculta == False,
            PlanoContas.is_deleted == False,
        )
        .order_by(PlanoContas.codigo)
    ).all()
    return contas


@router.get(
    "/exportar",
    dependencies=[Depends(require_any_permission(["plano_contas:view", "page:lancamentos:view", "page:dre:view"]))],
)
def exportar_plano_contas(
    *,
    db: Session = Depends(get_db),
    empresa_id: int = Depends(get_empresa_id_from_user),
):
    """Exporta o plano de contas com código e nome em CSV."""
    contas = db.exec(
        select(PlanoContas)
        .where(
            PlanoContas.empresa_id == empresa_id,
            PlanoContas.oculta == False,
            PlanoContas.is_deleted == False,
        )
        .order_by(PlanoContas.codigo)
    ).all()

    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')
    writer.writerow(["codigo", "nome"])
    for conta in contas:
        writer.writerow([conta.codigo or "", conta.nome or ""])

    csv_bytes = io.BytesIO(output.getvalue().encode("utf-8-sig"))
    csv_bytes.seek(0)

    return StreamingResponse(
        csv_bytes,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="plano_de_contas.csv"'},
    )


@router.post(
    "/bulk",
    response_model=List[PlanoContasRead],
    status_code=201,
    dependencies=[Depends(require_permission("plano_contas:create"))],
)
def create_plano_contas_bulk(
    *,
    db: Session = Depends(get_db),
    items_in: List[PlanoContasBulkCreateItem],
    empresa_id: int = Depends(get_empresa_id_from_user),
):
    """Cria categorias em massa, evitando duplicidades por nome/tipo/pai."""
    if not items_in:
        return []

    contas_empresa = db.exec(
        select(PlanoContas).where(
            PlanoContas.empresa_id == empresa_id,
            PlanoContas.oculta == False,
            PlanoContas.is_deleted == False,
        )
    ).all()

    ids_validos = {int(item.id) for item in contas_empresa if item.id is not None}
    existing_by_key: Dict[tuple[str, str, Optional[int]], PlanoContas] = {}
    for conta in contas_empresa:
        key = (
            _normalizar_tipo_plano(conta.tipo),
            _normalizar_chave_nome(conta.nome),
            int(conta.conta_pai_id) if conta.conta_pai_id is not None else None,
        )
        existing_by_key[key] = conta

    resolved: List[PlanoContas] = []
    for item in items_in:
        nome = str(item.nome or "").strip()
        if not nome:
            continue

        parent_id = int(item.conta_pai_id) if item.conta_pai_id is not None else None
        if parent_id is not None and parent_id not in ids_validos:
            raise HTTPException(status_code=400, detail=f"Categoria pai {parent_id} nao encontrada na empresa")

        tipo = _normalizar_tipo_plano(item.tipo)
        key = (tipo, _normalizar_chave_nome(nome), parent_id)
        existing = existing_by_key.get(key)
        if existing:
            resolved.append(existing)
            continue

        codigo = _normalizar_codigo_hierarquia(item.codigo)
        conta = PlanoContas(
            nome=nome,
            tipo=tipo,
            codigo=codigo,
            empresa_id=empresa_id,
            conta_pai_id=parent_id,
            permite_lancamentos=bool(item.permite_lancamentos if item.permite_lancamentos is not None else True),
            eh_operacional=bool(item.eh_operacional if item.eh_operacional is not None else True),
            considerar_nos_resultados=True,
            dre_grupo=crud_plano_contas._normalizar_dre_grupo(item.dre_grupo, tipo),
            oculta=False,
        )
        db.add(conta)
        db.flush()
        if conta.id is None:
            raise HTTPException(status_code=400, detail="Falha ao criar categoria no bulk")

        ids_validos.add(int(conta.id))
        existing_by_key[key] = conta
        resolved.append(conta)

    db.commit()

    uniq_ids: set[int] = set()
    uniq_resolved: List[PlanoContas] = []
    for conta in resolved:
        if conta.id is None:
            continue
        conta_id = int(conta.id)
        if conta_id in uniq_ids:
            continue
        uniq_ids.add(conta_id)
        db.refresh(conta)
        uniq_resolved.append(conta)

    crud_plano_contas.normalize_company_operational_categories(db=db, empresa_id=empresa_id)
    crud_plano_contas.sync_company_operational_hierarchy(db=db, empresa_id=empresa_id)
    return uniq_resolved


@router.post(
    "/bulk-sync",
    status_code=200,
    dependencies=[Depends(require_permission("plano_contas:update"))],
)
def bulk_sync_plano_contas(
    *,
    db: Session = Depends(get_db),
    payload: PlanoContasBulkSyncPayload,
    empresa_id: int = Depends(get_empresa_id_from_user),
):
    """Sincroniza todo o plano de contas em uma unica operacao (create/update/delete/reordenacao)."""
    if not payload.items:
        raise HTTPException(status_code=400, detail="Nenhuma categoria recebida para sincronizacao")

    try:
        contas_db = db.exec(
            select(PlanoContas).where(
                PlanoContas.empresa_id == empresa_id,
                PlanoContas.oculta == False,
                PlanoContas.is_deleted == False,
            )
        ).all()
        existing_by_id = {int(conta.id): conta for conta in contas_db if conta.id is not None}
        existing_ids = set(existing_by_id.keys())

        positive_ids = {int(item.id) for item in payload.items if int(item.id) > 0}
        unknown_ids = sorted(positive_ids - existing_ids)
        if unknown_ids:
            raise HTTPException(status_code=400, detail=f"IDs de categorias invalidos no payload: {unknown_ids}")

        temp_to_real: Dict[int, int] = {}
        resolved_by_payload_id: Dict[int, PlanoContas] = {}
        created_count = 0
        updated_count = 0
        changed_tipo: Dict[int, str] = {}

        for item in payload.items:
            payload_id = int(item.id)
            if payload_id > 0:
                resolved_by_payload_id[payload_id] = existing_by_id[payload_id]
                continue

            nome = str(item.nome or "").strip()
            if not nome:
                raise HTTPException(status_code=400, detail="Categoria nova sem nome no payload")

            tipo = _normalizar_tipo_plano(item.tipo)
            codigo = _normalizar_codigo_hierarquia(item.codigo)

            created = PlanoContas(
                nome=nome,
                tipo=tipo,
                codigo=codigo,
                empresa_id=empresa_id,
                conta_pai_id=None,
                permite_lancamentos=bool(item.permite_lancamentos if item.permite_lancamentos is not None else True),
                eh_operacional=bool(item.eh_operacional if item.eh_operacional is not None else True),
                considerar_nos_resultados=True,
                dre_grupo=crud_plano_contas._normalizar_dre_grupo(item.dre_grupo, tipo),
                oculta=False,
            )
            db.add(created)
            db.flush()
            if created.id is None:
                raise HTTPException(status_code=400, detail="Falha ao criar categoria nova no bulk-sync")

            created_count += 1
            temp_to_real[payload_id] = int(created.id)
            resolved_by_payload_id[payload_id] = created

        def resolve_parent_id(raw_parent_id: Optional[int]) -> Optional[int]:
            if raw_parent_id is None:
                return None
            parent_id = int(raw_parent_id)
            if parent_id > 0:
                if parent_id not in existing_ids and parent_id not in set(temp_to_real.values()):
                    raise HTTPException(status_code=400, detail=f"Categoria pai {parent_id} nao encontrada")
                return parent_id

            resolved = temp_to_real.get(parent_id)
            if resolved is None:
                raise HTTPException(status_code=400, detail=f"Categoria pai temporaria {parent_id} nao resolvida")
            return resolved

        final_parent_map: Dict[int, Optional[int]] = {}
        for item in payload.items:
            payload_id = int(item.id)
            conta = resolved_by_payload_id[payload_id]
            if conta.id is None:
                raise HTTPException(status_code=400, detail="Categoria sem ID durante sincronizacao")

            conta_id = int(conta.id)
            parent_id = resolve_parent_id(item.conta_pai_id)
            if parent_id == conta_id:
                raise HTTPException(status_code=400, detail=f"Categoria {conta_id} nao pode ser pai de si mesma")

            old_snapshot = (
                conta.nome,
                conta.tipo,
                conta.codigo,
                conta.conta_pai_id,
                conta.permite_lancamentos,
                conta.eh_operacional,
                conta.dre_grupo,
            )
            old_tipo = conta.tipo

            conta.nome = str(item.nome or "").strip()
            if not conta.nome:
                raise HTTPException(status_code=400, detail=f"Categoria {conta_id} sem nome")

            conta.tipo = _normalizar_tipo_plano(item.tipo)
            conta.codigo = _normalizar_codigo_hierarquia(item.codigo)
            conta.conta_pai_id = parent_id
            conta.permite_lancamentos = bool(item.permite_lancamentos if item.permite_lancamentos is not None else conta.permite_lancamentos)
            conta.eh_operacional = bool(item.eh_operacional if item.eh_operacional is not None else conta.eh_operacional)
            conta.considerar_nos_resultados = True
            conta.dre_grupo = crud_plano_contas._normalizar_dre_grupo(item.dre_grupo, conta.tipo)
            conta.oculta = False

            final_parent_map[conta_id] = parent_id
            db.add(conta)

            new_snapshot = (
                conta.nome,
                conta.tipo,
                conta.codigo,
                conta.conta_pai_id,
                conta.permite_lancamentos,
                conta.eh_operacional,
                conta.dre_grupo,
            )

            if payload_id > 0 and new_snapshot != old_snapshot:
                updated_count += 1

            if old_tipo != conta.tipo:
                changed_tipo[conta_id] = conta.tipo

        for conta_id in final_parent_map:
            seen: set[int] = set()
            current = conta_id
            while current is not None:
                if current in seen:
                    raise HTTPException(status_code=400, detail="Hierarquia circular detectada no plano de contas")
                seen.add(current)
                current = final_parent_map.get(current)

        deleted_ids = sorted(
            existing_ids - positive_ids,
            key=lambda current_id: len(str(existing_by_id[current_id].codigo or "").split(".")),
            reverse=True,
        )

        if deleted_ids:
            usage = db.exec(
                select(Lancamento.plano_contas_id, func.count(Lancamento.id))
                .where(
                    Lancamento.empresa_id == empresa_id,
                    Lancamento.plano_contas_id.in_(deleted_ids),
                    Lancamento.is_deleted == False,
                )
                .group_by(Lancamento.plano_contas_id)
            ).all()
            if usage:
                first_id = int(usage[0][0])
                first_count = int(usage[0][1])
                categoria = existing_by_id.get(first_id)
                nome = categoria.nome if categoria else str(first_id)
                raise HTTPException(
                    status_code=400,
                    detail=f"Nao foi possivel excluir '{nome}' pois existem {first_count} lancamento(s) vinculados",
                )

            usage_cartao = db.exec(
                select(LancamentoCartao.plano_contas_id)
                .where(
                    LancamentoCartao.empresa_id == empresa_id,
                    LancamentoCartao.plano_contas_id.in_(deleted_ids),
                    LancamentoCartao.is_deleted == False,
                )
            ).first()
            if usage_cartao:
                categoria = existing_by_id.get(int(usage_cartao))
                nome = categoria.nome if categoria else str(usage_cartao)
                raise HTTPException(
                    status_code=400,
                    detail=f"Não foi possível excluir '{nome}' pois existem despesas de cartão vinculadas.",
                )

            for conta_id in deleted_ids:
                item = existing_by_id[conta_id]
                item.is_deleted = True
                item.deleted_at = datetime.utcnow()
                db.add(item)

        if changed_tipo:
            lancamentos = db.exec(
                select(Lancamento).where(
                    Lancamento.empresa_id == empresa_id,
                    Lancamento.plano_contas_id.in_(list(changed_tipo.keys())),
                )
            ).all()
            for lanc in lancamentos:
                plano_id = int(lanc.plano_contas_id)
                lanc.tipo = "RECEITA" if changed_tipo.get(plano_id) == "R" else "DESPESA"
                db.add(lanc)

        db.commit()
        crud_plano_contas.normalize_company_operational_categories(db=db, empresa_id=empresa_id)
        crud_plano_contas.sync_company_operational_hierarchy(db=db, empresa_id=empresa_id)

        return {
            "created": created_count,
            "updated": updated_count,
            "deleted": len(deleted_ids),
            "total": len(payload.items),
        }

    except HTTPException:
        db.rollback()
        raise
    except Exception as exc:
        db.rollback()
        logger.error(f"Erro no bulk-sync do plano de contas: {exc}")
        raise HTTPException(status_code=400, detail="Erro ao sincronizar plano de contas em massa")


@router.post(
    "/importar",
    status_code=200,
    dependencies=[Depends(require_permission("plano_contas:create"))],
)
def importar_plano_contas_xlsx(
    *,
    db: Session = Depends(get_db),
    file: UploadFile = File(...),
    empresa_id: int = Depends(get_empresa_id_from_user),
):
    """Importa plano de contas via XLSX/CSV com inferencia de categoria pai por codigo hierarquico."""
    filename = (file.filename or "").strip()
    if not filename:
        raise HTTPException(status_code=400, detail="Arquivo nao informado")

    lowered = filename.lower()
    if not (lowered.endswith(".xlsx") or lowered.endswith(".xlsm") or lowered.endswith(".csv")):
        raise HTTPException(status_code=400, detail="Formato invalido. Envie um arquivo .xlsx, .xlsm ou .csv")

    content = file.file.read(MAX_PLANO_FILE_SIZE + 1)
    if len(content) > MAX_PLANO_FILE_SIZE:
        raise HTTPException(status_code=413, detail="Arquivo muito grande. Máximo permitido: 10MB.")
    if not content:
        raise HTTPException(status_code=400, detail="Arquivo vazio")

    rows: list[list[Any]] = []
    if lowered.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="ignore")
        reader = csv.reader(io.StringIO(text), delimiter=';')
        rows = [list(row) for row in reader]
    else:
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Nao foi possivel ler o arquivo XLSX: {exc}")

        sheet = workbook[workbook.sheetnames[0]]
        for row in sheet.iter_rows(values_only=True):
            rows.append(list(row))

    if len(rows) < 2:
        raise HTTPException(status_code=400, detail="Arquivo sem dados para importacao")

    header = [_normalizar_chave_nome(str(value or "")) for value in rows[0]]
    idx_classificacao = _pick_header_index(header, ["CLASSIFICACAO", "CLASSIFICAO"])
    idx_codigo = _pick_header_index(header, ["CODIGO", "COD", "CONTA", "CLASSIFICACAO"])
    idx_nome = _pick_header_index(header, ["NOME", "DESCRICAO", "CLASSIFICACAO", "CLASSIFICAO"])
    idx_tipo = _pick_header_index(header, ["TIPO", "NATUREZA"])
    idx_parent_code = _pick_header_index(header, ["CONTA_PAI_CODIGO", "CODIGO_PAI", "PARENT_CODE"])
    idx_endividamento = _pick_header_index(header, ["ENDIVIDAMENTO", "DIVIDA", "DIVIDA?"])
    idx_operacional = _pick_header_index(header, ["OPERACIONAL", "EH_OPERACIONAL", "OP", "OPERACIONAL?"])

    parsed_rows: list[dict[str, Any]] = []
    skipped_rows = 0
    for row_number, row in enumerate(rows[1:], start=2):
        raw_classificacao = row[idx_classificacao] if idx_classificacao is not None and idx_classificacao < len(row) else None
        raw_codigo = row[idx_codigo] if idx_codigo is not None and idx_codigo < len(row) else None
        raw_nome = row[idx_nome] if idx_nome is not None and idx_nome < len(row) else None
        raw_tipo = row[idx_tipo] if idx_tipo is not None and idx_tipo < len(row) else None
        raw_parent_code = row[idx_parent_code] if idx_parent_code is not None and idx_parent_code < len(row) else None
        raw_endividamento = row[idx_endividamento] if idx_endividamento is not None and idx_endividamento < len(row) else None
        raw_operacional = row[idx_operacional] if idx_operacional is not None and idx_operacional < len(row) else None

        codigo_from_classificacao, nome_from_classificacao = _parse_classificacao(raw_classificacao)
        codigo = _normalizar_codigo_hierarquia(raw_codigo) or codigo_from_classificacao
        nome = str(raw_nome or nome_from_classificacao or "").strip()
        if codigo is None and nome:
            codigo = _normalizar_codigo_hierarquia(raw_classificacao)

        if not nome:
            skipped_rows += 1
            continue
        if not codigo:
            skipped_rows += 1
            continue

        eh_operacional = True
        if raw_operacional is not None and str(raw_operacional).strip():
            eh_operacional = _is_true_like(raw_operacional)
        elif raw_endividamento is not None and str(raw_endividamento).strip():
            eh_operacional = not bool(_is_true_like(raw_endividamento))

        parsed_rows.append(
            {
                "line": row_number,
                "codigo": codigo,
                "nome": nome,
                "tipo": _normalizar_tipo_planilha(raw_tipo),
                "parent_code": _normalizar_codigo_hierarquia(raw_parent_code) if raw_parent_code is not None else None,
                "eh_operacional": eh_operacional,
            }
        )

    if not parsed_rows:
        raise HTTPException(status_code=400, detail="Nenhuma linha valida foi encontrada no arquivo")

    by_code: Dict[str, dict[str, Any]] = {}
    for item in parsed_rows:
        by_code[item["codigo"]] = item

    entries = list(by_code.values())
    explicit_tipo = {item["codigo"]: item["tipo"] for item in entries if item.get("tipo")}

    def infer_tipo(code: str) -> str:
        if code in explicit_tipo:
            return explicit_tipo[code]

        parts = code.split(".")
        for idx in range(len(parts) - 1, 0, -1):
            parent_code = ".".join(parts[:idx])
            if parent_code in explicit_tipo:
                return explicit_tipo[parent_code]

        descendants = [
            tipo for candidate_code, tipo in explicit_tipo.items()
            if candidate_code.startswith(f"{code}.")
        ]
        if descendants:
            counter = Counter(descendants)
            return counter.most_common(1)[0][0]

        return "R" if code.startswith("01") else "D"

    children_codes: set[str] = set()
    for item in entries:
        parts = item["codigo"].split(".")
        if len(parts) > 1:
            children_codes.add(".".join(parts[:-1]))

    entries.sort(key=lambda item: (len(item["codigo"].split(".")), item["codigo"]))

    contas_db = db.exec(
        select(PlanoContas).where(
            PlanoContas.empresa_id == empresa_id,
            PlanoContas.oculta == False,
        )
    ).all()

    existing_by_code: Dict[str, PlanoContas] = {}
    existing_by_name_tipo: Dict[tuple[str, str], List[PlanoContas]] = {}
    for conta in contas_db:
        norm_code = _normalizar_codigo_hierarquia(conta.codigo)
        if norm_code:
            existing_by_code[norm_code] = conta
        key = (_normalizar_chave_nome(conta.nome), _normalizar_tipo_plano(conta.tipo))
        existing_by_name_tipo.setdefault(key, []).append(conta)

    used_existing_ids: set[int] = set()
    code_to_real_id: Dict[str, int] = {}
    created_count = 0
    updated_count = 0

    for item in entries:
        codigo = item["codigo"]
        nome = str(item["nome"]).strip()
        tipo = infer_tipo(codigo)

        parent_code = item.get("parent_code")
        if not parent_code:
            parts = codigo.split(".")
            parent_code = ".".join(parts[:-1]) if len(parts) > 1 else None
        parent_id = code_to_real_id.get(parent_code) if parent_code else None

        has_children = codigo in children_codes
        permite_lancamentos = not has_children

        target = existing_by_code.get(codigo)
        if target and target.id is not None:
            used_existing_ids.add(int(target.id))

        if target is None:
            key = (_normalizar_chave_nome(nome), tipo)
            candidates = existing_by_name_tipo.get(key, [])
            for candidate in candidates:
                if candidate.id is None:
                    continue
                candidate_id = int(candidate.id)
                if candidate_id in used_existing_ids:
                    continue
                target = candidate
                used_existing_ids.add(candidate_id)
                break

        if target is None:
            target = PlanoContas(
                nome=nome,
                tipo=tipo,
                codigo=codigo,
                empresa_id=empresa_id,
                conta_pai_id=parent_id,
                permite_lancamentos=permite_lancamentos,
                eh_operacional=bool(item.get("eh_operacional", True)),
                considerar_nos_resultados=True,
                dre_grupo=crud_plano_contas._normalizar_dre_grupo(None, tipo),
                oculta=False,
            )
            db.add(target)
            db.flush()
            if target.id is None:
                raise HTTPException(status_code=400, detail="Falha ao criar categoria durante importacao")
            created_count += 1
        else:
            old_snapshot = (
                target.nome,
                target.tipo,
                target.codigo,
                target.conta_pai_id,
                target.permite_lancamentos,
                target.eh_operacional,
                target.dre_grupo,
            )
            target.nome = nome
            target.tipo = tipo
            target.codigo = codigo
            target.conta_pai_id = parent_id
            target.permite_lancamentos = permite_lancamentos
            target.eh_operacional = bool(item.get("eh_operacional", True))
            target.considerar_nos_resultados = True
            target.dre_grupo = crud_plano_contas._normalizar_dre_grupo(target.dre_grupo, tipo)
            db.add(target)

            new_snapshot = (
                target.nome,
                target.tipo,
                target.codigo,
                target.conta_pai_id,
                target.permite_lancamentos,
                target.eh_operacional,
                target.dre_grupo,
            )
            if old_snapshot != new_snapshot:
                updated_count += 1

        target_id = int(target.id)
        code_to_real_id[codigo] = target_id
        existing_by_code[codigo] = target

    db.commit()
    crud_plano_contas.normalize_company_operational_categories(db=db, empresa_id=empresa_id)
    crud_plano_contas.sync_company_operational_hierarchy(db=db, empresa_id=empresa_id)

    return {
        "arquivo": filename,
        "total_linhas": len(entries),
        "categorias_criadas": created_count,
        "categorias_atualizadas": updated_count,
        "linhas_ignoradas": skipped_rows,
    }

@router.post("", response_model=PlanoContasRead, status_code=201, include_in_schema=False, dependencies=[Depends(require_permission("plano_contas:create"))])
@router.post(
    "/",
    response_model=PlanoContasRead,
    status_code=201,
    dependencies=[Depends(require_permission("plano_contas:create"))],
)
def create_plano_contas(
    *,
    db: Session = Depends(get_db),
    conta_in: PlanoContasCreate,
    empresa_id: int = Depends(get_empresa_id_from_user),
    current_user: Usuario = Depends(get_current_active_user),
):
    """Cria uma nova categoria no plano de contas."""
    logger.info(f"Empresa {empresa_id} criando categoria: '{conta_in.nome}'")
    if conta_in.conta_pai_id is not None:
        pai = db.exec(
            select(PlanoContas).where(
                PlanoContas.id == conta_in.conta_pai_id,
                PlanoContas.empresa_id == empresa_id,
                PlanoContas.oculta == False,
                PlanoContas.is_deleted == False,
            )
        ).first()
        if not pai:
            raise HTTPException(status_code=400, detail="Categoria pai inválida ou não pertencente a esta empresa.")

    conta = crud_plano_contas.create(db=db, obj_in=conta_in, empresa_id=empresa_id)
    if conta_in.conta_pai_id is not None or not crud_plano_contas.can_manage_operational_flag(current_user.email):
        crud_plano_contas.sync_company_operational_hierarchy(db=db, empresa_id=empresa_id)
        conta = crud_plano_contas.get(db=db, id=int(conta.id or 0), empresa_id=empresa_id) or conta
    logger.success(f"Categoria '{conta.nome}' criada com ID: {conta.id}")
    return conta

@router.post(
    "/reordenar",
    status_code=200,
    dependencies=[Depends(require_permission("plano_contas:reorder"))],
)
def reordenar_plano_contas(
    *,
    db: Session = Depends(get_db),
    itens: List[ReordenacaoItem],
    empresa_id: int = Depends(get_empresa_id_from_user),
    current_user: Usuario = Depends(get_current_active_user),
):
    """
    Recebe a estrutura completa (Drag & Drop) e salva códigos e hierarquia em massa.
    Resolve o bug de 'perder números' ao recarregar.
    """
    try:
        # Prepara IDs para busca rápida
        ids = [item.id for item in itens]
        
        # Busca todas as contas envolvidas
        db_contas = db.exec(
            select(PlanoContas).where(
                PlanoContas.id.in_(ids),
                PlanoContas.empresa_id == empresa_id,
                PlanoContas.is_deleted == False,
            )
        ).all()
        
        # Mapa { ID: ObjetoBanco }
        conta_map = {c.id: c for c in db_contas}

        # Validar que todos os conta_pai_id referenciados nos itens pertencem à mesma empresa
        parent_ids = {int(item.conta_pai_id) for item in itens if item.conta_pai_id is not None}
        if parent_ids:
            valid_parent_ids = set(
                db.exec(
                    select(PlanoContas.id).where(
                        PlanoContas.id.in_(list(parent_ids)),
                        PlanoContas.empresa_id == empresa_id,
                        PlanoContas.is_deleted == False,
                    )
                ).all()
            )
            invalid_parents = parent_ids - valid_parent_ids
            if invalid_parents:
                raise HTTPException(
                    status_code=400,
                    detail=f"Categoria(s) pai {list(invalid_parents)} não pertencem a esta empresa."
                )
        
        updates = 0
        for item in itens:
            if item.id in conta_map:
                conta_db = conta_map[item.id]
                novo_tipo = _normalizar_tipo_plano(item.tipo)
                
                # Só atualiza se mudou algo (performance)
                if (conta_db.codigo != item.codigo or 
                    conta_db.conta_pai_id != item.conta_pai_id or
                    conta_db.tipo != novo_tipo):
                    tipo_alterado = conta_db.tipo != novo_tipo
                    
                    conta_db.codigo = item.codigo
                    conta_db.conta_pai_id = item.conta_pai_id
                    conta_db.tipo = novo_tipo
                    
                    db.add(conta_db)

                    if tipo_alterado:
                        novo_tipo_lanc = "RECEITA" if novo_tipo == "R" else "DESPESA"
                        lancamentos_categoria = db.exec(
                            select(Lancamento).where(
                                Lancamento.empresa_id == empresa_id,
                                Lancamento.plano_contas_id == conta_db.id
                            )
                        ).all()
                        for lanc in lancamentos_categoria:
                            lanc.tipo = novo_tipo_lanc
                            db.add(lanc)

                    updates += 1
        
        db.commit()
        crud_plano_contas.sync_company_operational_hierarchy(db=db, empresa_id=empresa_id)
        logger.info(f"Reordenação concluída. {updates} categorias atualizadas.")
        return {"message": "Ordem salva com sucesso"}
        
    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        logger.error(f"Erro ao reordenar: {e}")
        raise HTTPException(status_code=400, detail="Erro ao salvar a nova ordem.")

@router.patch(
    "/{conta_id}",
    response_model=PlanoContasRead,
    dependencies=[Depends(require_permission("plano_contas:update"))],
)
def update_plano_contas(
    *,
    db: Session = Depends(get_db),
    conta_id: int,
    conta_in: PlanoContasUpdate,
    empresa_id: int = Depends(get_empresa_id_from_user),
    current_user: Usuario = Depends(get_current_active_user),
):
    """
    Atualiza uma categoria. 
    SEGURANÇA: Bloqueia mudança de CÓDIGO se já existirem lançamentos.
    """
    logger.info(f"Empresa {empresa_id} atualizando categoria ID: {conta_id}")
    
    # 1. Busca a conta existente
    db_obj = crud_plano_contas.get(db=db, id=conta_id, empresa_id=empresa_id)
    if not db_obj:
        raise HTTPException(status_code=404, detail="Categoria não encontrada")
    if db_obj.oculta:
        raise HTTPException(status_code=400, detail="Categoria técnica do sistema não pode ser alterada")

    if conta_in.conta_pai_id is not None:
        if conta_in.conta_pai_id == conta_id:
            raise HTTPException(status_code=400, detail="Categoria não pode ser pai de si mesma.")
        pai = db.exec(
            select(PlanoContas).where(
                PlanoContas.id == conta_in.conta_pai_id,
                PlanoContas.empresa_id == empresa_id,
                PlanoContas.oculta == False,
                PlanoContas.is_deleted == False,
            )
        ).first()
        if not pai:
            raise HTTPException(status_code=400, detail="Categoria pai inválida ou não pertencente a esta empresa.")

    # 2. VERIFICAÇÃO DE SEGURANÇA (Se tentar mudar o código)
    if conta_in.codigo is not None and conta_in.codigo != db_obj.codigo:
        # Verifica uso em lançamentos
        uso = db.exec(
            select(func.count(Lancamento.id))
            .where(
                Lancamento.plano_contas_id == conta_id,
                Lancamento.is_deleted == False
            )
        ).one()

        if uso > 0:
            raise HTTPException(
                status_code=400, 
                detail=f"Não é permitido alterar a posição/código desta categoria pois existem {uso} lançamentos vinculados. Apenas o NOME pode ser editado."
            )
    
    # 3. Atualiza
    conta = crud_plano_contas.update(db=db, db_obj=db_obj, obj_in=conta_in)
    if conta_in.conta_pai_id is not None or not crud_plano_contas.can_manage_operational_flag(current_user.email):
        crud_plano_contas.sync_company_operational_hierarchy(db=db, empresa_id=empresa_id)
        conta = crud_plano_contas.get(db=db, id=conta_id, empresa_id=empresa_id) or conta
    logger.success(f"Categoria ID {conta.id} atualizada com sucesso.")
    return conta

@router.delete(
    "/{conta_id}",
    dependencies=[Depends(require_permission("plano_contas:delete"))],
)
def delete_plano_contas(
    *,
    db: Session = Depends(get_db),
    conta_id: int,
    empresa_id: int = Depends(get_empresa_id_from_user),
):
    """Remove uma categoria via soft delete, protegendo integridade referencial."""
    
    # 1. Busca a categoria ativa
    conta = crud_plano_contas.get(db=db, id=conta_id, empresa_id=empresa_id)
    if not conta:
        raise HTTPException(status_code=404, detail="Categoria não encontrada")
    if conta.oculta:
        raise HTTPException(status_code=400, detail="Categoria técnica do sistema não pode ser excluída")

    # 2. Verifica se tem subcategorias ativas
    filhos = db.exec(
        select(PlanoContas.id).where(
            PlanoContas.empresa_id == empresa_id,
            PlanoContas.conta_pai_id == conta_id,
            PlanoContas.is_deleted == False,
        )
    ).first()
    if filhos:
        raise HTTPException(status_code=400, detail="Não é possível excluir uma categoria que possui subcategorias.")

    # 3. Verifica se tem lançamentos financeiros ativos
    uso = db.exec(
        select(Lancamento.id).where(
            Lancamento.empresa_id == empresa_id,
            Lancamento.plano_contas_id == conta_id,
            Lancamento.is_deleted == False,
        )
    ).first()
    if uso:
        raise HTTPException(status_code=400, detail="Não é possível excluir uma categoria que possui lançamentos.")

    # 4. Verifica se tem despesas de cartão vinculadas
    uso_cartao = db.exec(
        select(LancamentoCartao.id).where(
            LancamentoCartao.empresa_id == empresa_id,
            LancamentoCartao.plano_contas_id == conta_id,
            LancamentoCartao.is_deleted == False,
        )
    ).first()
    if uso_cartao:
        raise HTTPException(status_code=400, detail="Não é possível excluir uma categoria vinculada a despesas de cartão.")

    # 5. Verifica se é categoria padrão em integração bancária
    uso_integracao = db.exec(
        select(IntegracaoBancaria.id).where(
            IntegracaoBancaria.empresa_id == empresa_id,
            IntegracaoBancaria.categoria_padrao_id == conta_id,
            IntegracaoBancaria.is_deleted == False,
        )
    ).first()
    if uso_integracao:
        raise HTTPException(status_code=400, detail="Não é possível excluir a categoria padrão de uma integração bancária ativa.")

    # 6. Verifica regras de conciliação bancária (MapeamentoCategoria)
    uso_mapeamento = db.exec(
        select(MapeamentoCategoria.id)
        .join(IntegracaoBancaria, IntegracaoBancaria.id == MapeamentoCategoria.integracao_id)
        .where(
            IntegracaoBancaria.empresa_id == empresa_id,
            MapeamentoCategoria.plano_contas_id == conta_id,
            MapeamentoCategoria.is_deleted == False,
            IntegracaoBancaria.is_deleted == False,
        )
    ).first()
    if uso_mapeamento:
        raise HTTPException(status_code=400, detail="Não é possível excluir uma categoria vinculada a regras de conciliação bancária.")

    db_obj = crud_plano_contas.delete(db=db, id=conta_id, empresa_id=empresa_id)
    if not db_obj:
        raise HTTPException(status_code=404, detail="Categoria não encontrada")
    
    logger.success(f"Categoria ID {conta_id} removida com sucesso.")
    return {"ok": True}
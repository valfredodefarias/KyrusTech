# app/services/importacao_service.py
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime
from decimal import Decimal
import io
import os
import re
from typing import Any, Callable, Iterator, Optional, cast
import uuid
import zipfile
import pandas as pd
from openpyxl import load_workbook
from fastapi import HTTPException, status
from loguru import logger
from sqlmodel import Session, select

from app.db.session import engine
from app.models.import_job import ImportJob
from app.models.conta import Conta
from app.models.plano_contas import PlanoContas
from app.models.centro_custo import CentroCusto
from app.models.entidade import Entidade
from app.models.lancamento import Lancamento
from app.crud import crud_plano_contas
from app.services.aprendizado_service import (
    normalizar_texto_importacao,
    load_learning_references,
    build_learning_reference_index,
    build_inference_cache_parallel,
    build_import_suggestions_from_samples,
    infer_learning_ids,
    infer_learning_ids_cached,
)

IMPORT_ANALYZE_SAMPLE_LIMIT = 24
IMPORT_INSERT_BATCH_SIZE = 5000
IMPORT_PREPARE_CHUNK_SIZE = 2000
IMPORT_MAX_WORKERS = max(2, min(4, os.cpu_count() or 2))


def create_import_job(db: Session, kind: str, empresa_id: int, user_id: int, filename: str) -> ImportJob:
    job = ImportJob(
        job_id=str(uuid.uuid4()),
        kind=kind,
        empresa_id=empresa_id,
        user_id=user_id,
        filename=filename,
    )
    db.add(job)
    db.commit()
    db.refresh(job)
    return job


def update_import_job(
    job_id: str,
    *,
    status: Optional[str] = None,
    progress: Optional[int] = None,
    message: Optional[str] = None,
    error: Optional[str] = None,
    result: Optional[dict[str, Any]] = None,
) -> None:
    def _apply_update(session: Session) -> None:
        job = session.get(ImportJob, job_id)
        if not job:
            return
        if status is not None:
            job.status = status
        if progress is not None:
            job.progress = max(0, min(100, int(progress)))
        if message is not None:
            job.message = message
        if error is not None:
            job.error = error
        if result is not None:
            job.result = result
        job.updated_at = datetime.utcnow()
        session.add(job)
        session.commit()

    with Session(engine) as session:
        _apply_update(session)


def serialize_import_job(job: ImportJob) -> dict[str, Any]:
    payload = job.dict()
    payload["created_at"] = job.created_at.isoformat() if job.created_at else None
    payload["updated_at"] = job.updated_at.isoformat() if job.updated_at else None
    return payload


def load_import_system_rows(session: Session, empresa_id: int) -> dict[str, list[dict[str, Any]]]:
    crud_plano_contas.normalize_company_operational_categories(db=session, empresa_id=empresa_id)
    contas = session.exec(
        select(Conta.id, Conta.nome).where(Conta.empresa_id == empresa_id)
    ).all()
    categorias = session.exec(select(PlanoContas).where(PlanoContas.empresa_id == empresa_id)).all()
    centros = session.exec(
        select(CentroCusto.id, CentroCusto.nome).where(CentroCusto.empresa_id == empresa_id)
    ).all()
    entidades = session.exec(
        select(Entidade.id, Entidade.nome).where(Entidade.empresa_id == empresa_id)
    ).all()

    return {
        "contas": [{"id": conta_id, "nome": nome} for conta_id, nome in contas if conta_id is not None],
        "categorias": [
            {
                "id": categoria.id,
                "nome": categoria.nome,
                "tipo": categoria.tipo,
                "codigo": categoria.codigo,
                "conta_pai_id": categoria.conta_pai_id,
                "permite_lancamentos": categoria.permite_lancamentos,
                "eh_cabecalho": categoria.eh_cabecalho,
            }
            for categoria in categorias
            if categoria.id is not None
        ],
        "centros": [{"id": centro_id, "nome": nome} for centro_id, nome in centros if centro_id is not None],
        "entidades": [{"id": entidade_id, "nome": nome} for entidade_id, nome in entidades if entidade_id is not None],
    }


def ensure_import_prerequisites(sistema: dict[str, list[dict[str, Any]]]) -> None:
    sem_contas = len(sistema.get("contas") or []) == 0
    sem_entidades = len(sistema.get("entidades") or []) == 0

    if sem_contas and sem_entidades:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Antes de importar, cadastre pelo menos 1 banco em Contas Bancarias e pelo menos 1 interessado em Interessados.",
        )

    if sem_contas:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Antes de importar, cadastre pelo menos 1 banco em Contas Bancarias.",
        )

    if sem_entidades:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Antes de importar, cadastre pelo menos 1 interessado em Interessados.",
        )


def normalize_import_lookup_key(value: Any) -> str:
    import unicodedata
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    normalized = normalized.replace("\u00a0", " ").replace("\u200b", "").replace("\ufeff", "")
    normalized = re.sub(r"\s+", " ", normalized).strip().upper()
    return normalized


def build_import_name_map(rows: list[dict[str, Any]], selectable_only: bool = False) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for row in rows:
        if selectable_only and not (row.get("permite_lancamentos", True) and not row.get("eh_cabecalho", False)):
            continue
        if row.get("id") is None:
            continue
        key = normalize_import_lookup_key(row.get("nome"))
        if not key:
            continue
        mapping[key] = int(row["id"])
    return mapping


def extract_import_category_code(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    match = re.match(r"^\s*(\d+(?:\s*[\.\-]\s*\d+)+)", text)
    if not match:
        return ""
    parts = re.findall(r"\d+", match.group(1))
    return ".".join(parts)


def build_import_category_code_map(rows: list[dict[str, Any]], selectable_only: bool = False) -> dict[str, int]:
    code_to_ids: dict[str, set[int]] = defaultdict(set)
    for row in rows:
        if selectable_only and not (row.get("permite_lancamentos", True) and not row.get("eh_cabecalho", False)):
            continue
        if row.get("id") is None:
            continue
        row_id = int(row["id"])
        code = extract_import_category_code(row.get("codigo")) or extract_import_category_code(row.get("nome"))
        if not code:
            continue
        code_to_ids[code].add(row_id)
    return {code: next(iter(ids)) for code, ids in code_to_ids.items() if len(ids) == 1}


def resolve_import_category_alias(
    cat_key: str,
    nomes_cats_sist: dict[str, int],
    codigos_cats_sist: dict[str, int],
) -> Optional[int]:
    if not cat_key:
        return None

    if cat_key == "SALDO INICIAL":
        alias_candidates = [
            "05.05 RECEBIMENTO DE EMPRESTIMOS A SOCIOS",
            "RECEBIMENTO DE EMPRESTIMOS A SOCIOS",
            "RECEBIMENTO DE EMPRESTIMOS A SOCIO",
        ]
        for alias in alias_candidates:
            alias_key = normalize_import_lookup_key(alias)
            if alias_key in nomes_cats_sist:
                return int(nomes_cats_sist[alias_key])

        if "05.05" in codigos_cats_sist:
            return int(codigos_cats_sist["05.05"])

        semantic_matches = {
            int(cat_id)
            for nome, cat_id in nomes_cats_sist.items()
            if "EMPRESTIM" in nome and "SOCI" in nome and ("RECEB" in nome or "ENTRADA" in nome)
        }
        if len(semantic_matches) == 1:
            return next(iter(semantic_matches))

    return None


def format_preview_value(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if pd.isna(value):
        return ""
    if isinstance(value, (date, datetime)):
        try:
            return value.strftime("%Y-%m-%d")
        except Exception:
            return str(value)
    return str(value).strip()


def normalize_import_headers(header_values: tuple[Any, ...] | list[Any]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(header_values):
        header = str(value or f"COLUNA_{index + 1}").upper().strip() or f"COLUNA_{index + 1}"
        if header in seen:
            seen[header] += 1
            header = f"{header}_{seen[header]}"
        else:
            seen[header] = 1
        headers.append(header)
    return headers


def iter_spreadsheet_rows(file_bytes: bytes) -> tuple[list[str], Iterator[dict[str, str]]]:
    buffer = io.BytesIO(file_bytes)
    if zipfile.is_zipfile(buffer):
        workbook = load_workbook(buffer, read_only=True, data_only=True)
        worksheet = workbook.worksheets[0]
        row_iter = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(row_iter)
        except StopIteration:
            workbook.close()
            return [], iter(())

        headers = normalize_import_headers(header_row or [])

        def generator() -> Iterator[dict[str, str]]:
            try:
                for row in row_iter:
                    values = list(row or [])
                    padded = values + [""] * max(0, len(headers) - len(values))
                    yield {
                        headers[index]: format_preview_value(padded[index] if index < len(padded) else "")
                        for index in range(len(headers))
                    }
            finally:
                workbook.close()

        return headers, generator()

    dataframe = pd.read_excel(io.BytesIO(file_bytes), sheet_name=0, dtype=str, keep_default_na=False)
    dataframe.columns = [str(column).upper().strip() for column in dataframe.columns]
    dataframe = dataframe.fillna("")
    headers = [str(column).upper().strip() for column in dataframe.columns]

    def fallback_generator() -> Iterator[dict[str, str]]:
        for row in dataframe.to_dict(orient="records"):
            yield {str(key).upper().strip(): format_preview_value(value) for key, value in row.items()}

    return headers, fallback_generator()


def find_column_in_headers(headers: list[str], possiveis_nomes: list[str]) -> str:
    header_map = {str(header).upper().strip(): str(header).upper().strip() for header in headers}
    for nome in possiveis_nomes:
        normalized = str(nome).upper().strip()
        if normalized in header_map:
            return header_map[normalized]
    return ""


def coerce_row_value(row: dict[str, Any], column: str) -> str:
    if not column:
        return ""
    value = row.get(column, "")
    return format_preview_value(value)


def resolve_import_value_raw(
    row: dict[str, Any],
    col_valor: str,
    col_valor_pago: str,
    col_valor_previsto: str,
    col_data_pagamento: str,
) -> str:
    if col_valor:
        return coerce_row_value(row, col_valor)

    valor_pago = coerce_row_value(row, col_valor_pago)
    valor_previsto = coerce_row_value(row, col_valor_previsto)
    data_pagamento = coerce_row_value(row, col_data_pagamento)

    if data_pagamento:
        return valor_pago or valor_previsto
    return valor_previsto or valor_pago


def parse_import_decimal(raw_value: Any, cache: dict[str, Decimal]) -> Decimal:
    key = format_preview_value(raw_value)
    if key in cache:
        return cache[key]
    cleaned = key.replace("R$", "").replace(" ", "")
    if "," in cleaned and "." in cleaned:
        if cleaned.rfind(",") > cleaned.rfind("."):
            cleaned = cleaned.replace(".", "").replace(",", ".")
        else:
            cleaned = cleaned.replace(",", "")
    elif "," in cleaned:
        cleaned = cleaned.replace(",", ".")
    cleaned = cleaned or "0"
    cache[key] = Decimal(cleaned)
    return cache[key]


def build_import_duplicate_signature(
    data_vencimento: date,
    descricao_normalizada: str,
    tipo: str,
    valor: Decimal,
    conta_id: Optional[int],
) -> tuple[date, str, str, str, Optional[int]]:
    valor_norm = Decimal(str(valor)).quantize(Decimal("0.01"))
    return (
        data_vencimento,
        descricao_normalizada,
        str(tipo or "").upper().strip(),
        f"{valor_norm:.2f}",
        int(conta_id) if conta_id is not None else None,
    )


def split_import_date_parts(raw_value: str) -> Optional[tuple[int, int, int]]:
    cleaned = str(raw_value or "").strip()
    if not cleaned:
        return None

    match = re.match(r"^(\d{1,4})[\/-](\d{1,2})[\/-](\d{1,4})$", cleaned)
    if not match:
        return None

    left, middle, right = (int(part) for part in match.groups())
    if left >= 1000 or right >= 1000:
        return left, middle, right
    return None


def infer_import_dayfirst(date_values: list[Any]) -> bool:
    dayfirst_votes = 0
    monthfirst_votes = 0

    for raw_value in date_values:
        key = format_preview_value(raw_value)
        if not key:
            continue

        parts = split_import_date_parts(key)
        if not parts:
            continue

        left, middle, right = parts
        if left >= 1000 or right < 1000:
            continue
        if left > 12 and middle <= 12:
            dayfirst_votes += 1
        elif middle > 12 and left <= 12:
            monthfirst_votes += 1

    return dayfirst_votes >= monthfirst_votes


def parse_import_date(raw_value: Any, cache: dict[str, Optional[date]], dayfirst: bool = True) -> Optional[date]:
    key = format_preview_value(raw_value)
    cache_key = f"{'DMY' if dayfirst else 'MDY'}::{key}"
    if cache_key in cache:
        return cache[cache_key]
    if not key:
        cache[cache_key] = None
        return None

    if isinstance(raw_value, datetime):
        cache[cache_key] = raw_value.date()
        return cache[cache_key]
    if isinstance(raw_value, date):
        cache[cache_key] = raw_value
        return cache[cache_key]

    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M"):
        try:
            cache[cache_key] = datetime.strptime(key, fmt).date()
            return cache[cache_key]
        except ValueError:
            pass

    parts = split_import_date_parts(key)
    if parts:
        left, middle, right = parts
        try:
            if left >= 1000:
                cache[cache_key] = date(left, middle, right)
            elif right >= 1000:
                day = left if dayfirst else middle
                month = middle if dayfirst else left
                cache[cache_key] = date(right, month, day)
            else:
                cache[cache_key] = None
            return cache[cache_key]
        except ValueError:
            cache[cache_key] = None
            return None

    parsed = pd.to_datetime(key, dayfirst=dayfirst, errors="coerce")
    cache[cache_key] = parsed.date() if parsed is not None and not pd.isna(parsed) else None
    return cache[cache_key]


def prepare_import_chunk(
    raw_chunk: list[dict[str, Any]],
    map_categorias: dict[str, Any],
    map_contas: dict[str, Any],
    map_centros: dict[str, Any],
    map_entidades: dict[str, Any],
    nomes_cats_sist: dict[str, int],
    codigos_cats_sist: dict[str, int],
    nomes_contas_sist: dict[str, int],
    nomes_centros_sist: dict[str, int],
    nomes_entidades_sist: dict[str, int],
    cache_tipos: dict[int, Any],
    inference_cache: dict[tuple[str, str], dict[str, Optional[int] | float]],
) -> tuple[list[dict[str, Any]], dict[str, str], dict[str, str]]:
    prepared_rows: list[dict[str, Any]] = []
    missing_entities: dict[str, str] = {}
    missing_centers: dict[str, str] = {}

    for raw in raw_chunk:
        descricao = str(raw["descricao"] or "")
        tipo = str(raw["tipo"] or "")
        categoria_nome = str(raw["categoria_nome"] or "")
        entidade_nome = str(raw["entidade_nome"] or "")
        conta_nome = str(raw["conta_nome"] or "")
        centro_nome = str(raw["centro_nome"] or "")
        inferencia = inference_cache.get((descricao.strip(), tipo.strip()), {"plano_contas_id": None, "entidade_id": None}) if descricao else {"plano_contas_id": None, "entidade_id": None}

        plano_contas_id: Optional[int] = None
        cat_key = normalize_import_lookup_key(categoria_nome)
        cat_code = extract_import_category_code(categoria_nome)
        cat_alias_id = resolve_import_category_alias(cat_key, nomes_cats_sist, codigos_cats_sist)
        if cat_key and cat_key in map_categorias:
            plano_contas_id = int(map_categorias[cat_key])
        elif cat_key and cat_key in nomes_cats_sist:
            plano_contas_id = int(nomes_cats_sist[cat_key])
        elif cat_code and cat_code in codigos_cats_sist:
            plano_contas_id = int(codigos_cats_sist[cat_code])
        elif cat_alias_id is not None:
            plano_contas_id = int(cat_alias_id)
        elif inferencia.get("plano_contas_id") is not None:
            plano_contas_id = int(cast(int, inferencia.get("plano_contas_id")))

        if not tipo and plano_contas_id is not None:
            tipo = "RECEITA" if cache_tipos.get(int(plano_contas_id)) == "R" else "DESPESA"
        if not tipo:
            tipo = "DESPESA"

        entidade_id: Optional[int] = None
        ent_key = normalize_import_lookup_key(entidade_nome)
        if ent_key and ent_key in map_entidades:
            entidade_id = int(map_entidades[ent_key])
        elif ent_key and ent_key in nomes_entidades_sist:
            entidade_id = int(nomes_entidades_sist[ent_key])
        elif inferencia.get("entidade_id") is not None:
            entidade_id = int(cast(int, inferencia.get("entidade_id")))
        elif ent_key:
            missing_entities[ent_key] = entidade_nome

        conta_id: Optional[int] = None
        conta_key = normalize_import_lookup_key(conta_nome)
        if conta_key:
            if conta_key in map_contas:
                conta_id = int(map_contas[conta_key])
            elif conta_key in nomes_contas_sist:
                conta_id = int(nomes_contas_sist[conta_key])

        centro_custo_id: Optional[int] = None
        centro_key = normalize_import_lookup_key(centro_nome)
        if centro_key:
            if centro_key in map_centros:
                centro_custo_id = int(map_centros[centro_key])
            elif centro_key in nomes_centros_sist:
                centro_custo_id = int(nomes_centros_sist[centro_key])
            else:
                missing_centers[centro_key] = centro_nome

        prepared_rows.append(
            {
                **raw,
                "tipo": tipo,
                "plano_contas_id": plano_contas_id,
                "entidade_id": entidade_id,
                "conta_id": conta_id,
                "centro_custo_id": centro_custo_id,
                "entidade_key": ent_key,
                "centro_key": centro_key,
            }
        )

    return prepared_rows, missing_entities, missing_centers


def flush_lancamento_batch(db: Session, batch: list[Lancamento]) -> None:
    if not batch:
        return
    db.add_all(batch)
    db.flush()
    for lancamento in batch:
        db.expunge(lancamento)


def analyze_import_contents(session: Session, file_bytes: bytes, empresa_id: int, progress_callback: Optional[Callable[[int, str], None]] = None) -> dict[str, Any]:
    if progress_callback:
        progress_callback(5, "Lendo arquivo XLSX em streaming")
    headers, row_iter = iter_spreadsheet_rows(file_bytes)

    col_desc = find_column_in_headers(headers, ["DESCRIÇÃO", "DESCRICAO", "HISTÓRICO", "HISTORICO"])
    col_tipo = find_column_in_headers(headers, ["TIPO"])
    col_venc = find_column_in_headers(headers, ["DATA VENCIMENTO", "VENCIMENTO", "DATA"])
    col_pag = find_column_in_headers(headers, ["DATA PAGAMENTO", "PAGAMENTO"])
    col_valor = find_column_in_headers(headers, ["VALOR"])
    col_valor_pago = find_column_in_headers(headers, ["VALOR PAGO"])
    col_valor_previsto = find_column_in_headers(headers, ["VALOR PREVISTO"])
    col_conta = find_column_in_headers(headers, ["CONTA", "BANCO"])
    col_cat = find_column_in_headers(headers, ["CATEGORIA", "PLANO DE CONTAS"])
    col_centro = find_column_in_headers(headers, ["CENTRO DE CUSTO", "CENTRO", "FILIAL", "CENTRO_CUSTO"])
    col_entidade = find_column_in_headers(headers, ["ENTIDADE", "CLIENTE", "FORNECEDOR"])

    if not col_valor and not col_valor_pago and not col_valor_previsto:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Arquivo sem coluna de valor. Informe VALOR ou VALOR PAGO/VALOR PREVISTO.",
        )

    sistema = load_import_system_rows(session, empresa_id)
    ensure_import_prerequisites(sistema)
    nomes_contas = build_import_name_map(sistema["contas"])
    nomes_cats = build_import_name_map(sistema["categorias"], selectable_only=True)
    codigos_cats = build_import_category_code_map(sistema["categorias"], selectable_only=True)
    nomes_centros = build_import_name_map(sistema["centros"])
    nomes_entidades = build_import_name_map(sistema["entidades"])

    preview_source: list[dict[str, Any]] = []
    conflitos_contas: set[str] = set()
    conflitos_categorias: set[str] = set()
    conflitos_centros: set[str] = set()
    conflitos_entidades: set[str] = set()
    categoria_samples: dict[str, dict[tuple[str, str], int]] = defaultdict(lambda: defaultdict(int))
    entidade_samples: dict[str, dict[tuple[str, str], int]] = defaultdict(lambda: defaultdict(int))
    inference_keys: set[tuple[str, str]] = set()
    total_rows = 0

    for row_idx, row in enumerate(row_iter, start=2):
        total_rows += 1
        if progress_callback and total_rows % 1000 == 0:
            progress_callback(min(40, 8 + min(32, total_rows // 1000)), f"Indexando planilha ({total_rows} linhas lidas)")

        descricao = coerce_row_value(row, col_desc)
        tipo_raw = coerce_row_value(row, col_tipo).upper()
        tipo = "RECEITA" if tipo_raw.startswith("R") else ("DESPESA" if tipo_raw.startswith("D") else "")
        categoria_nome = coerce_row_value(row, col_cat)
        entidade_nome = coerce_row_value(row, col_entidade)
        conta_nome = coerce_row_value(row, col_conta)
        centro_nome = coerce_row_value(row, col_centro)
        categoria_key = normalize_import_lookup_key(categoria_nome)
        categoria_code = extract_import_category_code(categoria_nome)
        categoria_alias_id = resolve_import_category_alias(categoria_key, nomes_cats, codigos_cats)
        entidade_key = normalize_import_lookup_key(entidade_nome)
        conta_key = normalize_import_lookup_key(conta_nome)
        centro_key = normalize_import_lookup_key(centro_nome)

        if conta_key and conta_key not in nomes_contas:
            conflitos_contas.add(conta_nome)
        if categoria_key and categoria_key not in nomes_cats and not (categoria_code and categoria_code in codigos_cats) and categoria_alias_id is None:
            conflitos_categorias.add(categoria_nome)
            if descricao:
                categoria_samples[categoria_nome][(descricao, tipo)] += 1
                inference_keys.add((descricao.strip(), tipo.strip()))
        if centro_key and centro_key not in nomes_centros:
            conflitos_centros.add(centro_nome)
        if entidade_key and entidade_key not in nomes_entidades:
            conflitos_entidades.add(entidade_nome)
            if descricao:
                entidade_samples[entidade_nome][(descricao, tipo)] += 1
                inference_keys.add((descricao.strip(), tipo.strip()))

        preview_source.append(
            {
                "linha": row_idx,
                "descricao": descricao,
                "tipo": tipo if tipo else coerce_row_value(row, col_tipo),
                "valor": resolve_import_value_raw(
                    row,
                    col_valor,
                    col_valor_pago,
                    col_valor_previsto,
                    col_pag,
                ),
                "data_vencimento": coerce_row_value(row, col_venc),
                "categoria_arquivo": categoria_nome,
                "entidade_arquivo": entidade_nome,
            }
        )

    inferred_dayfirst = infer_import_dayfirst([item.get("data_vencimento") for item in preview_source])

    conflitos = {
        "contas": sorted(conflitos_contas),
        "categorias": sorted(conflitos_categorias),
        "centros": sorted(conflitos_centros),
        "entidades": sorted(conflitos_entidades),
    }

    if progress_callback:
        progress_callback(55, "Calculando sugestões por amostragem")
    learning_refs = load_learning_references(session, empresa_id)
    learning_ref_index = build_learning_reference_index(learning_refs)
    inference_cache = build_inference_cache_parallel(inference_keys, learning_refs, learning_ref_index)
    sugestoes = build_import_suggestions_from_samples(
        categoria_samples,
        entidade_samples,
        sistema["categorias"],
        learning_refs,
        inference_cache,
        learning_ref_index,
    )
    categorias_by_id = {int(item["id"]): item for item in sistema["categorias"] if item.get("id") is not None}
    entidades_by_id = {int(item["id"]): item for item in sistema["entidades"] if item.get("id") is not None}
    preview_rows: list[dict[str, Any]] = []

    if progress_callback:
        progress_callback(75, "Montando prévia das primeiras linhas")
    for row in preview_source:
        descricao = str(row.get("descricao") or "")
        tipo = str(row.get("tipo") or "")
        categoria_nome = str(row.get("categoria_arquivo") or "")
        entidade_nome = str(row.get("entidade_arquivo") or "")

        categoria_key = normalize_import_lookup_key(categoria_nome)
        categoria_code = extract_import_category_code(categoria_nome)
        categoria_alias_id = resolve_import_category_alias(categoria_key, nomes_cats, codigos_cats)
        entidade_key = normalize_import_lookup_key(entidade_nome)
        categoria_mapeada_id = nomes_cats.get(categoria_key) if categoria_key else None
        if categoria_mapeada_id is None and categoria_code:
            categoria_mapeada_id = codigos_cats.get(categoria_code)
        if categoria_mapeada_id is None and categoria_alias_id is not None:
            categoria_mapeada_id = categoria_alias_id
        entidade_mapeada_id = nomes_entidades.get(entidade_key) if entidade_key else None

        if categoria_mapeada_id is None and categoria_nome:
            categoria_mapeada_id = sugestoes["categorias"].get(categoria_nome)
        if entidade_mapeada_id is None and entidade_nome:
            entidade_mapeada_id = sugestoes["entidades"].get(entidade_nome)

        if descricao and (categoria_mapeada_id is None or entidade_mapeada_id is None):
            inferencia = infer_learning_ids_cached(descricao, tipo, learning_refs, inference_cache, learning_ref_index)
            if categoria_mapeada_id is None and inferencia.get("plano_contas_id") is not None:
                categoria_mapeada_id = int(cast(int, inferencia.get("plano_contas_id")))
            if entidade_mapeada_id is None and inferencia.get("entidade_id") is not None:
                entidade_mapeada_id = int(cast(int, inferencia.get("entidade_id")))

        preview_rows.append(
            {
                **row,
                "categoria_sugerida_id": categoria_mapeada_id,
                "categoria_sugerida_nome": categorias_by_id.get(int(categoria_mapeada_id), {}).get("nome") if categoria_mapeada_id is not None else None,
                "entidade_sugerida_id": entidade_mapeada_id,
                "entidade_sugerida_nome": entidades_by_id.get(int(entidade_mapeada_id), {}).get("nome") if entidade_mapeada_id is not None else None,
            }
        )

    return {
        "conflitos": conflitos,
        "sistema": sistema,
        "sugestoes": sugestoes,
        "preview": preview_rows,
        "meta": {
            "total_linhas": total_rows,
            "linhas_preview": len(preview_rows),
            "date_format_detected": "DD/MM/YYYY" if inferred_dayfirst else "MM/DD/YYYY",
        },
    }


def run_import_analysis_job(job_id: str, empresa_id: int, user_id: int, file_bytes: bytes) -> None:
    try:
        update_import_job(job_id, status="RUNNING", progress=1, message="Iniciando análise do arquivo")
        with Session(engine) as session:
            session.info["audit_user_id"] = user_id
            result = analyze_import_contents(
                session,
                file_bytes,
                empresa_id,
                progress_callback=lambda progress, message: update_import_job(job_id, status="RUNNING", progress=progress, message=message),
            )
        update_import_job(job_id, status="COMPLETED", progress=100, message="Análise concluída", result=result)
    except Exception as exc:
        logger.exception("Erro no job de análise de importação")
        update_import_job(job_id, status="ERROR", progress=100, message="Falha na análise", error=str(exc))


def execute_import_contents(
    db: Session,
    file_bytes: bytes,
    empresa_id: int,
    user_id: int,
    mapeamento: dict[str, Any],
    progress_callback: Optional[Callable[[int, str], None]] = None,
) -> dict[str, Any]:
    if progress_callback:
        progress_callback(5, "Lendo arquivo XLSX em streaming")
    headers, row_iter = iter_spreadsheet_rows(file_bytes)

    map_categorias = {
        normalize_import_lookup_key(k): v
        for k, v in mapeamento.get("map_categorias", {}).items()
        if v is not None and str(v).strip() != ""
    }
    map_contas = {
        normalize_import_lookup_key(k): v
        for k, v in mapeamento.get("map_contas", {}).items()
        if v is not None and str(v).strip() != ""
    }
    map_centros = {
        normalize_import_lookup_key(k): v
        for k, v in mapeamento.get("map_centros", {}).items()
        if v is not None and str(v).strip() != ""
    }
    map_entidades = {
        normalize_import_lookup_key(k): v
        for k, v in mapeamento.get("map_entidades", {}).items()
        if v is not None and str(v).strip() != ""
    }

    col_venc = find_column_in_headers(headers, ["DATA VENCIMENTO", "VENCIMENTO", "DATA"])
    col_pag = find_column_in_headers(headers, ["DATA PAGAMENTO", "PAGAMENTO"])
    col_desc = find_column_in_headers(headers, ["DESCRIÇÃO", "DESCRICAO", "HISTÓRICO", "HISTORICO"])
    col_valor = find_column_in_headers(headers, ["VALOR"])
    col_valor_pago = find_column_in_headers(headers, ["VALOR PAGO"])
    col_valor_previsto = find_column_in_headers(headers, ["VALOR PREVISTO"])
    col_cat = find_column_in_headers(headers, ["CATEGORIA", "PLANO DE CONTAS"])
    col_ent = find_column_in_headers(headers, ["ENTIDADE", "CLIENTE", "FORNECEDOR"])
    col_conta = find_column_in_headers(headers, ["CONTA", "BANCO"])
    col_centro = find_column_in_headers(headers, ["CENTRO DE CUSTO", "CENTRO", "FILIAL", "CENTRO_CUSTO"])
    col_tipo = find_column_in_headers(headers, ["TIPO"])

    if not col_valor and not col_valor_pago and not col_valor_previsto:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Arquivo sem coluna de valor. Informe VALOR ou VALOR PAGO/VALOR PREVISTO.",
        )

    sistema = load_import_system_rows(db, empresa_id)
    ensure_import_prerequisites(sistema)
    cache_tipos = {int(item["id"]): item.get("tipo") for item in sistema["categorias"] if item.get("id") is not None}
    nomes_cats_sist = build_import_name_map(sistema["categorias"], selectable_only=True)
    codigos_cats_sist = build_import_category_code_map(sistema["categorias"], selectable_only=True)
    nomes_contas_sist = build_import_name_map(sistema["contas"])
    nomes_centros_sist = build_import_name_map(sistema["centros"])
    nomes_entidades_sist = build_import_name_map(sistema["entidades"])
    learning_refs = load_learning_references(db, empresa_id)
    learning_ref_index = build_learning_reference_index(learning_refs)

    raw_rows: list[dict[str, Any]] = []
    inference_keys: set[tuple[str, str]] = set()
    parse_date_cache: dict[str, Optional[date]] = {}
    parse_decimal_cache: dict[str, Decimal] = {}

    for row_idx, row in enumerate(row_iter, start=2):
        descricao = coerce_row_value(row, col_desc)
        categoria_nome = coerce_row_value(row, col_cat)
        entidade_nome = coerce_row_value(row, col_ent)
        conta_nome = coerce_row_value(row, col_conta)
        centro_nome = coerce_row_value(row, col_centro)
        tipo_raw = coerce_row_value(row, col_tipo).upper()
        tipo = "RECEITA" if tipo_raw.startswith("R") else ("DESPESA" if tipo_raw.startswith("D") else "")
        raw_rows.append(
            {
                "linha": row_idx,
                "descricao": descricao,
                "categoria_nome": categoria_nome,
                "entidade_nome": entidade_nome,
                "conta_nome": conta_nome,
                "centro_nome": centro_nome,
                "tipo": tipo,
                "valor_raw": resolve_import_value_raw(
                    row,
                    col_valor,
                    col_valor_pago,
                    col_valor_previsto,
                    col_pag,
                ),
                "data_venc_raw": coerce_row_value(row, col_venc),
                "data_pag_raw": coerce_row_value(row, col_pag),
            }
        )
        if descricao:
            categoria_key = normalize_import_lookup_key(categoria_nome)
            categoria_code = extract_import_category_code(categoria_nome)
            categoria_alias_id = resolve_import_category_alias(categoria_key, nomes_cats_sist, codigos_cats_sist)
            entidade_key = normalize_import_lookup_key(entidade_nome)
            categoria_resolvida = (
                not categoria_key
                or categoria_key in map_categorias
                or categoria_key in nomes_cats_sist
                or (categoria_code and categoria_code in codigos_cats_sist)
                or categoria_alias_id is not None
            )
            entidade_resolvida = not entidade_key or entidade_key in map_entidades or entidade_key in nomes_entidades_sist
            if not categoria_resolvida or not entidade_resolvida:
                inference_keys.add((descricao.strip(), tipo.strip()))

    total_rows = len(raw_rows)
    inferred_dayfirst = infer_import_dayfirst([
        *(row.get("data_venc_raw") for row in raw_rows),
        *(row.get("data_pag_raw") for row in raw_rows),
    ])
    if progress_callback:
        progress_callback(18, f"Preparando {total_rows} linha(s) para importação")

    if progress_callback:
        progress_callback(24, "Calculando inferências reutilizáveis")
    inference_cache = build_inference_cache_parallel(inference_keys, learning_refs, learning_ref_index)

    missing_entities: dict[str, str] = {}
    missing_centers: dict[str, str] = {}
    prepared_rows: list[dict[str, Any]] = []
    raw_chunks = [raw_rows[index:index + IMPORT_PREPARE_CHUNK_SIZE] for index in range(0, total_rows, IMPORT_PREPARE_CHUNK_SIZE)]
    with ThreadPoolExecutor(max_workers=IMPORT_MAX_WORKERS) as executor:
        prepared_chunks = list(
            executor.map(
                lambda chunk: prepare_import_chunk(
                    chunk,
                    map_categorias,
                    map_contas,
                    map_centros,
                    map_entidades,
                    nomes_cats_sist,
                    codigos_cats_sist,
                    nomes_contas_sist,
                    nomes_centros_sist,
                    nomes_entidades_sist,
                    cache_tipos,
                    inference_cache,
                ),
                raw_chunks,
            )
        )

    for index, (chunk_rows, chunk_missing_entities, chunk_missing_centers) in enumerate(prepared_chunks, start=1):
        prepared_rows.extend(chunk_rows)
        missing_entities.update(chunk_missing_entities)
        missing_centers.update(chunk_missing_centers)
        if progress_callback:
            progress_callback(24 + int((index / max(len(prepared_chunks), 1)) * 21), f"Resolvendo mapeamentos por lote ({index}/{max(len(prepared_chunks), 1)})")

    uncategorized_rows = [row for row in prepared_rows if row.get("plano_contas_id") is None]
    if uncategorized_rows:
        linhas = ", ".join(str(row.get("linha")) for row in uncategorized_rows[:20])
        sufixo = "" if len(uncategorized_rows) <= 20 else ", ..."
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                f"Existem {len(uncategorized_rows)} lançamento(s) sem categoria definida. "
                f"Revise o mapeamento e categorize todas as linhas antes de confirmar. "
                f"Linhas: {linhas}{sufixo}"
            ),
        )

    if progress_callback:
        progress_callback(45, "Criando entidades e centros ausentes")
    if missing_entities:
        novos = [Entidade(nome=nome, tipo="AMBOS", cpf_cnpj=None, status="ATIVO", empresa_id=empresa_id) for nome in missing_entities.values()]
        db.add_all(novos)
        db.flush()
        for ent_key, entidade in zip(missing_entities.keys(), novos):
            nomes_entidades_sist[ent_key] = int(cast(int, entidade.id))

    if missing_centers:
        novos_centros = [CentroCusto(nome=nome, empresa_id=empresa_id) for nome in missing_centers.values()]
        db.add_all(novos_centros)
        db.flush()
        for centro_key, centro in zip(missing_centers.keys(), novos_centros):
            nomes_centros_sist[centro_key] = int(cast(int, centro.id))

    if progress_callback:
        progress_callback(55, "Inserindo lançamentos em lotes")
    erros: list[str] = []
    importados = 0
    ignorados_duplicidade = 0
    batch: list[Lancamento] = []

    existing_import_rows = db.exec(
        select(
            Lancamento.data_vencimento,
            Lancamento.descricao,
            Lancamento.tipo,
            Lancamento.valor_previsto,
            Lancamento.conta_id,
        ).where(
            Lancamento.empresa_id == empresa_id,
            Lancamento.is_deleted == False,
            Lancamento.origem == "IMPORTACAO",
        )
    ).all()
    existing_signature_counts: dict[tuple[date, str, str, str, Optional[int]], int] = defaultdict(int)
    for existing_date, existing_desc, existing_tipo, existing_valor, existing_conta_id in existing_import_rows:
        if existing_date is None:
            continue
        normalized_desc = normalizar_texto_importacao(existing_desc)
        if not normalized_desc:
            continue
        valor_assinatura = Decimal(str(existing_valor if existing_valor is not None else 0))
        signature = build_import_duplicate_signature(
            existing_date,
            normalized_desc,
            str(existing_tipo or ""),
            valor_assinatura,
            int(existing_conta_id) if existing_conta_id is not None else None,
        )
        existing_signature_counts[signature] += 1

    incoming_signature_counts: dict[tuple[date, str, str, str, Optional[int]], int] = defaultdict(int)

    for index, row in enumerate(prepared_rows, start=1):
        try:
            data_vencimento = parse_import_date(row["data_venc_raw"], parse_date_cache, dayfirst=inferred_dayfirst)
            if data_vencimento is None:
                continue
            data_pagamento = parse_import_date(row["data_pag_raw"], parse_date_cache, dayfirst=inferred_dayfirst)
            valor = parse_import_decimal(row["valor_raw"], parse_decimal_cache)
            plano_contas_id = row["plano_contas_id"]
            conta_id = int(row["conta_id"]) if row["conta_id"] else None

            descricao_normalizada = normalizar_texto_importacao(str(row["descricao"] or ""))
            if descricao_normalizada:
                signature = build_import_duplicate_signature(
                    data_vencimento,
                    descricao_normalizada,
                    str(row["tipo"]),
                    valor,
                    conta_id,
                )
                existing_count = existing_signature_counts[signature]
                seen_count = incoming_signature_counts[signature]
                incoming_signature_counts[signature] += 1
                if seen_count < existing_count:
                    ignorados_duplicidade += 1
                    continue

            status_lanc = "PAGO" if data_pagamento else "ABERTO"
            data_competencia = data_pagamento or data_vencimento
            mes_ano = f"{data_competencia.month:02d}-{data_competencia.year}"
            novo_lancamento = Lancamento(
                descricao=str(row["descricao"] or "Sem descrição"),
                tipo=str(row["tipo"]),
                status=status_lanc,
                valor_previsto=valor,
                valor_pago=valor if status_lanc == "PAGO" else None,
                data_vencimento=data_vencimento,
                data_pagamento=data_pagamento,
                data_competencia=data_competencia,
                competencia=mes_ano,
                plano_contas_id=plano_contas_id,
                entidade_id=int(row["entidade_id"]) if row["entidade_id"] else None,
                conta_id=conta_id,
                centro_custo_id=int(row["centro_custo_id"]) if row["centro_custo_id"] else None,
                empresa_id=empresa_id,
                created_by_id=user_id,
                origem="IMPORTACAO",
                observacao=f"Importado da linha {row['linha']}",
            )
            batch.append(novo_lancamento)
            importados += 1

            if len(batch) >= IMPORT_INSERT_BATCH_SIZE:
                flush_lancamento_batch(db, batch)
                batch.clear()

            if progress_callback and (index % 1000 == 0 or index == len(prepared_rows)):
                pct = 55 + int((index / max(len(prepared_rows), 1)) * 40)
                progress_callback(min(96, pct), f"Importando registros ({index}/{len(prepared_rows)})")
        except Exception as e:
            erros.append(f"Linha {row['linha']}: {str(e)}")

    if batch:
        flush_lancamento_batch(db, batch)
        batch.clear()

    db.commit()

    if progress_callback:
        progress_callback(100, "Importação concluída com sucesso")

    return {
        "status": "sucesso",
        "importados": importados,
        "ignorados_duplicidade": ignorados_duplicidade,
        "erros": erros[:50],
        "total_erros": len(erros),
    }


def run_import_execute_job(job_id: str, empresa_id: int, user_id: int, file_bytes: bytes, mapeamento: dict[str, Any]) -> None:
    try:
        update_import_job(job_id, status="RUNNING", progress=1, message="Preparando importação")
        with Session(engine) as session:
            session.info["audit_user_id"] = user_id
            result = execute_import_contents(
                session,
                file_bytes,
                empresa_id,
                user_id,
                mapeamento,
                progress_callback=lambda progress, message: update_import_job(job_id, status="RUNNING", progress=progress, message=message),
            )
        update_import_job(job_id, status="COMPLETED", progress=100, message="Importação concluída", result=result)
    except Exception as exc:
        logger.exception("Erro no job de execução de importação")
        update_import_job(job_id, status="ERROR", progress=100, message="Falha na importação", error=str(exc))

"""
Serviço de integração com Itaú via importação de arquivos XLSX.
Processa extratos (recebimentos) e relatórios de pagamentos (despesas).
"""
from typing import List, Dict, Optional, Tuple
from datetime import datetime, date
from decimal import Decimal
from loguru import logger
import re
import hashlib
import json
import unicodedata
from openpyxl import load_workbook
from io import BytesIO

from sqlmodel import Session, select, or_
from app.models.lancamento import Lancamento
from app.models.plano_contas import PlanoContas
from app.models.entidade import Entidade


def limpar_valor_monetario(valor_str: str) -> Decimal:
    """
    Limpa e converte string monetária para Decimal.
    Ex: "R$ 1.234,56" -> Decimal("1234.56")
    """
    if not valor_str or valor_str == "-":
        return Decimal("0")
    
    # Remove símbolos e espaços
    valor_limpo = str(valor_str).replace("R$", "").replace(" ", "").replace(".", "").replace(",", ".")
    
    try:
        return Decimal(valor_limpo)
    except:
        return Decimal("0")


def parsear_data_hora(data_str: str) -> Tuple[Optional[date], Optional[datetime]]:
    """
    Parseia data e hora em vários formatos.
    """
    if not data_str:
        return None, None

    if isinstance(data_str, datetime):
        return data_str.date(), data_str

    if isinstance(data_str, date):
        return data_str, None
    
    formatos_hora = [
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
    ]

    valor = str(data_str).strip()
    for fmt in formatos_hora:
        try:
            data_hora = datetime.strptime(valor, fmt)
            return data_hora.date(), data_hora
        except:
            continue

    formatos_data = [
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%d/%m/%y",
    ]
    
    for fmt in formatos_data:
        try:
            return datetime.strptime(valor, fmt).date(), None
        except:
            continue
    
    return None, None


def parsear_data(data_str: str) -> Optional[date]:
    """
    Parseia data em vários formatos.
    """
    data, _ = parsear_data_hora(data_str)
    return data


def _normalizar_texto(texto: Optional[str]) -> str:
    if not texto:
        return ""
    return re.sub(r"\s+", " ", str(texto).strip().lower())


def _normalizar_cabecalho(texto: Optional[str]) -> str:
    if not texto:
        return ""
    valor = unicodedata.normalize("NFKD", str(texto))
    valor = "".join([c for c in valor if not unicodedata.combining(c)])
    valor = re.sub(r"\s+", " ", valor)
    return valor.strip().lower()


def _encontrar_cabecalho(ws, pistas_list: List[List[str]], linhas_max: int = 40) -> Optional[int]:
    for idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if idx > linhas_max:
            break
        linha_texto = " ".join([_normalizar_cabecalho(cell) for cell in row if cell is not None])
        for pistas in pistas_list:
            if all(pista in linha_texto for pista in pistas):
                return idx
    return None


def _limpar_cpf_cnpj(cpf_cnpj: Optional[str]) -> str:
    return re.sub(r"[^0-9]", "", cpf_cnpj or "")


def _status_aberto_clause() -> tuple[str, ...]:
    return ("PENDENTE", "EM ABERTO")


def gerar_import_hash(lancamento: Dict, conta_id: Optional[int] = None) -> str:
    payload = {
        "origem": lancamento.get("origem"),
        "tipo": lancamento.get("tipo"),
        "data": str(lancamento.get("data") or ""),
        "data_hora": lancamento.get("data_hora"),
        "valor": str(lancamento.get("valor") or ""),
        "descricao": _normalizar_texto(lancamento.get("descricao")),
        "razao_social": _normalizar_texto(lancamento.get("razao_social")),
        "cpf_cnpj": _limpar_cpf_cnpj(lancamento.get("cpf_cnpj")),
        "referencia": _normalizar_texto(lancamento.get("referencia")),
        "linha_arquivo": lancamento.get("linha_arquivo"),
        "conta_id": conta_id or lancamento.get("conta_id"),
    }
    payload_str = json.dumps(payload, sort_keys=True, ensure_ascii=False)
    return hashlib.sha256(payload_str.encode("utf-8")).hexdigest()


def processar_extrato_itau(arquivo_bytes: bytes, empresa_id: int) -> List[Dict]:
    """
    Processa arquivo XLSX de extrato do Itaú (recebimentos).
    
    Retorna lista de lançamentos encontrados no extrato.
    """
    logger.info("Processando extrato do Itaú...")
    
    wb = load_workbook(BytesIO(arquivo_bytes), data_only=True)
    ws = wb.active
    
    lancamentos = []
    
    # Procura pela linha de cabeçalho
    linha_cabecalho = _encontrar_cabecalho(
        ws,
        [
            ["data", "lancamento"],
            ["data", "lançamento"],
            ["data", "valor"],
        ]
    )
    
    if not linha_cabecalho:
        raise ValueError("Não foi possível encontrar o cabeçalho do extrato")
    
    # Identifica índices das colunas
    cabecalho = list(ws.iter_rows(min_row=linha_cabecalho, max_row=linha_cabecalho, values_only=True))[0]
    
    idx_data = None
    idx_lancamento = None
    idx_razao_social = None
    idx_cpf_cnpj = None
    idx_valor = None
    idx_saldo = None
    
    for i, cell in enumerate(cabecalho):
        if not cell:
            continue
        cell_str = str(cell).upper()
        if "DATA" in cell_str and idx_data is None:
            idx_data = i
        elif "LANÇAMENTO" in cell_str or "LANCAMENTO" in cell_str:
            idx_lancamento = i
        elif "RAZÃO" in cell_str or "RAZAO" in cell_str:
            idx_razao_social = i
        elif "CPF" in cell_str or "CNPJ" in cell_str:
            idx_cpf_cnpj = i
        elif "VALOR" in cell_str and "SALDO" not in cell_str:
            idx_valor = i
        elif "SALDO" in cell_str:
            idx_saldo = i
    
    if idx_data is None or idx_lancamento is None or idx_valor is None:
        raise ValueError("Estrutura do arquivo não reconhecida. Colunas obrigatórias não encontradas.")
    
    # Processa linhas de dados
    for row_idx, row in enumerate(ws.iter_rows(min_row=linha_cabecalho + 1, values_only=True), linha_cabecalho + 1):
        if not any(row):  # Linha vazia
            continue
        
        # Pula linha de saldo anterior
        lancamento_str = str(row[idx_lancamento] or "").upper()
        if "SALDO ANTERIOR" in lancamento_str or "SALDO INICIAL" in lancamento_str:
            continue
        
        data_str = row[idx_data]
        descricao = str(row[idx_lancamento] or "").strip()
        razao_social = str(row[idx_razao_social] or "").strip() if idx_razao_social is not None else ""
        cpf_cnpj = str(row[idx_cpf_cnpj] or "").strip() if idx_cpf_cnpj is not None else ""
        valor_str = row[idx_valor]
        
        if not descricao or descricao == "None":
            continue
        
        data_lancamento, data_hora = parsear_data_hora(data_str)
        if not data_lancamento:
            logger.warning(f"Data inválida na linha {row_idx}: {data_str}")
            continue
        
        valor = limpar_valor_monetario(valor_str)
        if valor == 0:
            continue
        
        # Determina se é receita (valores positivos no extrato geralmente são recebimentos)
        tipo = "RECEITA"
        if valor < 0:
            tipo = "DESPESA"
            valor = abs(valor)
        
        # Para extratos, a data do lançamento é tanto vencimento quanto pagamento (já foi realizado)
        lancamento = {
            "data": data_lancamento,  # Data do lançamento (usada como padrão)
            "data_pagamento": data_lancamento.isoformat() if data_lancamento else None,
            "data_vencimento": data_lancamento.isoformat() if data_lancamento else None,  # Mesma data (já foi realizado)
            "data_hora": data_hora.isoformat() if data_hora else None,
            "descricao": descricao,
            "razao_social": razao_social,
            "cpf_cnpj": cpf_cnpj,
            "valor": valor,
            "valor_pago": valor,  # Para extratos, valor pago = valor (já foi realizado)
            "valor_previsto": valor,  # Valor previsto também é o mesmo
            "tipo": tipo,
            "origem": "ITAU_EXTRATO",
            "linha_arquivo": row_idx
        }
        
        lancamentos.append(lancamento)
    
    logger.success(f"Processados {len(lancamentos)} lançamentos do extrato")
    return lancamentos


def processar_relatorio_pagamentos_itau(arquivo_bytes: bytes, empresa_id: int) -> List[Dict]:
    """
    Processa arquivo XLSX de relatório de pagamentos do Itaú (despesas).
    
    Retorna lista de pagamentos encontrados no relatório.
    """
    logger.info("Processando relatório de pagamentos do Itaú...")
    
    wb = load_workbook(BytesIO(arquivo_bytes), data_only=True)
    ws = wb.active
    
    lancamentos = []
    
    # Procura pela linha de cabeçalho
    linha_cabecalho = _encontrar_cabecalho(
        ws,
        [
            ["data", "pagamento", "valor"],
            ["data", "pagamentos", "valor"],
            ["data do pagamento", "valor"],
            ["favorecido", "data", "valor"],
            ["beneficiario", "data", "valor"],
        ]
    )
    
    if not linha_cabecalho:
        raise ValueError("Não foi possível encontrar o cabeçalho do relatório de pagamentos")
    
    # Identifica índices das colunas
    cabecalho = list(ws.iter_rows(min_row=linha_cabecalho, max_row=linha_cabecalho, values_only=True))[0]
    
    idx_favorecido = None
    idx_cpf_cnpj = None
    idx_tipo_pagamento = None
    idx_data = None
    idx_data_vencimento = None
    idx_valor = None
    idx_referencia = None
    
    for i, cell in enumerate(cabecalho):
        if not cell:
            continue
        cell_str = _normalizar_cabecalho(cell).upper()
        if "FAVORECIDO" in cell_str or "BENEFICIÁRIO" in cell_str or "BENEFICIARIO" in cell_str:
            idx_favorecido = i
        elif "CPF" in cell_str or "CNPJ" in cell_str:
            idx_cpf_cnpj = i
        elif "TIPO" in cell_str and "PAGAMENTO" in cell_str:
            idx_tipo_pagamento = i
        elif "DATA" in cell_str and "PAGAMENTO" in cell_str:
            idx_data = i
        elif "DATA" in cell_str and ("VENCIMENTO" in cell_str or "VENC" in cell_str):
            idx_data_vencimento = i
        elif "VALOR" in cell_str:
            idx_valor = i
        elif "REFERÊNCIA" in cell_str or "REFERENCIA" in cell_str:
            idx_referencia = i
    
    if idx_favorecido is None or idx_data is None or idx_valor is None:
        raise ValueError("Estrutura do arquivo não reconhecida. Colunas obrigatórias não encontradas.")
    
    # Processa linhas de dados
    for row_idx, row in enumerate(ws.iter_rows(min_row=linha_cabecalho + 1, values_only=True), linha_cabecalho + 1):
        if not any(row):  # Linha vazia
            continue
        
        favorecido = str(row[idx_favorecido] or "").strip()
        cpf_cnpj = str(row[idx_cpf_cnpj] or "").strip() if idx_cpf_cnpj is not None else ""
        tipo_pagamento = str(row[idx_tipo_pagamento] or "").strip() if idx_tipo_pagamento is not None else ""
        data_str = row[idx_data]
        data_vencimento_str = row[idx_data_vencimento] if idx_data_vencimento is not None else None
        valor_str = row[idx_valor]
        referencia = str(row[idx_referencia] or "").strip() if idx_referencia is not None else ""
        
        if not favorecido or favorecido == "None":
            continue
        
        data_pagamento, data_hora = parsear_data_hora(data_str)
        if not data_pagamento:
            logger.warning(f"Data inválida na linha {row_idx}: {data_str}")
            continue
        
        # Tenta parsear data de vencimento se disponível
        data_vencimento = None
        if data_vencimento_str:
            data_vencimento = parsear_data(data_vencimento_str)
        
        valor = limpar_valor_monetario(valor_str)
        if valor == 0:
            continue
        
        # Monta descrição
        descricao = f"{tipo_pagamento} - {favorecido}" if tipo_pagamento else favorecido
        if referencia:
            descricao += f" ({referencia})"
        
        lancamento = {
            "data": data_pagamento,  # Data de pagamento (usada como padrão)
            "data_pagamento": data_pagamento.isoformat() if data_pagamento else None,
            "data_vencimento": data_vencimento.isoformat() if data_vencimento else data_pagamento.isoformat(),  # Usa data de pagamento como fallback
            "data_hora": data_hora.isoformat() if data_hora else None,
            "descricao": descricao,
            "razao_social": favorecido,
            "cpf_cnpj": cpf_cnpj,
            "referencia": referencia,
            "valor": valor,
            "valor_pago": valor,  # Para pagamentos, valor pago = valor
            "valor_previsto": valor,  # Valor previsto também é o mesmo
            "tipo": "DESPESA",
            "origem": "ITAU_PAGAMENTOS",
            "linha_arquivo": row_idx
        }
        
        lancamentos.append(lancamento)
    
    logger.success(f"Processados {len(lancamentos)} pagamentos do relatório")
    return lancamentos


def verificar_duplicata(
    db: Session,
    lancamento: Dict,
    empresa_id: int,
    conta_id: Optional[int] = None
) -> Optional[Lancamento]:
    """
    Verifica duplicidade usando fingerprint import_hash.
    Retorna o lançamento existente se encontrar, None caso contrário.
    """
    import_hash = lancamento.get("import_hash") or gerar_import_hash(lancamento, conta_id=conta_id)
    if not import_hash:
        return None

    return db.exec(
        select(Lancamento).where(
            Lancamento.empresa_id == empresa_id,
            Lancamento.import_hash == import_hash
        )
    ).first()


def verificar_duplicata_ofx_por_fallback(
    db: Session,
    lancamento: Dict,
    empresa_id: int,
    conta_id: Optional[int] = None,
) -> Optional[Lancamento]:
    """
    Fallback defensivo para OFX sem identificador estavel de transacao.

    Quando o banco nao envia id/fitid/reference, o parser gera um fallback textual.
    Nesse caso tentamos localizar um movimento ja importado pela mesma conta,
    mesma data, mesmo valor, mesmo tipo e descricao normalizada igual.
    """
    if str(lancamento.get("origem") or "").upper() != "OFX_EXTRATO":
        return None

    referencia = _normalizar_texto(lancamento.get("referencia"))
    if referencia and not referencia.startswith("fallback"):
        return None

    conta_resolvida = conta_id or lancamento.get("conta_id")
    if not conta_resolvida:
        return None

    descricao = _normalizar_texto(lancamento.get("descricao"))
    if not descricao:
        return None

    data_base = lancamento.get("data")
    if isinstance(data_base, datetime):
        data_base = data_base.date()
    elif isinstance(data_base, str):
        data_base, _ = parsear_data_hora(data_base)

    if not data_base:
        return None

    try:
        valor = Decimal(str(lancamento.get("valor") or "0"))
    except Exception:
        return None

    candidatos = db.exec(
        select(Lancamento).where(
            Lancamento.empresa_id == empresa_id,
            Lancamento.conta_id == conta_resolvida,
            Lancamento.origem == "OFX_EXTRATO",
            Lancamento.tipo == lancamento.get("tipo"),
            or_(
                Lancamento.data_pagamento == data_base,
                Lancamento.data_vencimento == data_base,
            ),
        )
    ).all()

    for candidato in candidatos:
        descricao_candidata = _normalizar_texto(candidato.descricao)
        valor_candidato = candidato.valor_pago if candidato.valor_pago not in (None, Decimal("0.00")) else candidato.valor_previsto
        if descricao_candidata == descricao and Decimal(str(valor_candidato or "0")) == valor:
            return candidato

    return None


def buscar_lancamento_previsto_mesmo_dia_valor(
    db: Session,
    lancamento: Dict,
    empresa_id: int,
    centro_custo_id: Optional[int] = None,
    tolerancia_valor: Optional[Decimal] = None,
    tolerancia_percentual: Optional[Decimal] = None,
) -> Optional[Lancamento]:
    """
    Busca lançamento previsto no mesmo dia e com mesmo valor.
    Se centro_custo_id for fornecido, busca apenas naquele centro de custo.
    Caso contrário, busca em geral (todos os centros de custo).
    """
    data_lancamento = lancamento["data"]
    valor = Decimal(str(lancamento["valor"]))
    margem = abs(valor) * tolerancia_percentual if tolerancia_percentual is not None else (tolerancia_valor or Decimal("1.00"))
    valor_min = valor - margem
    valor_max = valor + margem
    
    query = select(Lancamento).where(
        Lancamento.empresa_id == empresa_id,
        Lancamento.tipo == lancamento.get("tipo"),
        Lancamento.data_vencimento == data_lancamento,
        Lancamento.valor_previsto >= valor_min,
        Lancamento.valor_previsto <= valor_max,
        Lancamento.status.in_(_status_aberto_clause())
    )
    
    # Se tem centro de custo, filtra por ele
    if centro_custo_id:
        query = query.where(Lancamento.centro_custo_id == centro_custo_id)
    
    lancamento_previsto = db.exec(query).first()
    
    return lancamento_previsto


def buscar_lancamento_atrasado_mesmo_valor(
    db: Session,
    lancamento: Dict,
    empresa_id: int,
    centro_custo_id: Optional[int] = None,
    dias_tolerancia: int = 30,
    tolerancia_valor: Optional[Decimal] = None,
    tolerancia_percentual: Optional[Decimal] = None,
) -> List[Lancamento]:
    """
    Busca lançamentos em atraso com mesmo valor.
    Se centro_custo_id for fornecido, busca apenas naquele centro de custo.
    Caso contrário, busca em geral (todos os centros de custo).
    """
    from datetime import timedelta
    
    data_lancamento = lancamento["data"]
    valor = Decimal(str(lancamento["valor"]))
    margem = abs(valor) * tolerancia_percentual if tolerancia_percentual is not None else (tolerancia_valor or Decimal("1.00"))
    valor_min = valor - margem
    valor_max = valor + margem
    data_limite = data_lancamento - timedelta(days=dias_tolerancia)
    
    query = select(Lancamento).where(
        Lancamento.empresa_id == empresa_id,
        Lancamento.tipo == lancamento.get("tipo"),
        Lancamento.data_vencimento < data_lancamento,
        Lancamento.data_vencimento >= data_limite,
        Lancamento.valor_previsto >= valor_min,
        Lancamento.valor_previsto <= valor_max,
        Lancamento.status.in_(_status_aberto_clause())
    )
    
    # Se tem centro de custo, filtra por ele
    if centro_custo_id:
        query = query.where(Lancamento.centro_custo_id == centro_custo_id)
    
    lancamentos_atrasados = db.exec(
        query.order_by(Lancamento.data_vencimento.desc())
    ).all()
    
    return list(lancamentos_atrasados)


def criar_entidade_se_nao_existir(
    db: Session,
    razao_social: str,
    cpf_cnpj: str,
    empresa_id: int
) -> Optional[int]:
    """
    Cria entidade se não existir, retorna ID.
    """
    if not razao_social and not cpf_cnpj:
        return None
    
    # Limpa CPF/CNPJ
    cpf_cnpj_limpo = re.sub(r'[^0-9]', '', cpf_cnpj) if cpf_cnpj else ""
    
    # Busca entidade existente
    if cpf_cnpj_limpo:
        entidade = db.exec(
            select(Entidade).where(
                Entidade.empresa_id == empresa_id,
                Entidade.cpf_cnpj == cpf_cnpj_limpo
            )
        ).first()
        
        if entidade:
            return entidade.id
    
    # Busca por nome
    if razao_social:
        entidade = db.exec(
            select(Entidade).where(
                Entidade.empresa_id == empresa_id,
                Entidade.nome.ilike(f"%{razao_social[:50]}%")
            )
        ).first()
        
        if entidade:
            return entidade.id
    
    # Cria nova entidade
    if razao_social or cpf_cnpj_limpo:
        tipo_pessoa = "PJ" if len(cpf_cnpj_limpo) > 11 else "PF"
        
        nova_entidade = Entidade(
            nome=razao_social or cpf_cnpj_limpo,
            tipo="AMBOS",
            tipo_pessoa=tipo_pessoa,
            cpf_cnpj=cpf_cnpj_limpo if cpf_cnpj_limpo else None,
            status="ATIVO",
            empresa_id=empresa_id
        )
        db.add(nova_entidade)
        db.commit()
        db.refresh(nova_entidade)
        
        logger.info(f"Entidade criada: {nova_entidade.nome} (ID: {nova_entidade.id})")
        return nova_entidade.id
    
    return None


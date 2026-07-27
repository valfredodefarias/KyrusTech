# scripts/align_all_imported_items.py
import os
import sys
import openpyxl
from pathlib import Path
from datetime import datetime, date
from decimal import Decimal
from sqlmodel import Session, select
from sqlalchemy import text

ROOT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT_DIR))

from app.db.session import engine
from app.models.lancamento import Lancamento
from app.models.plano_contas import PlanoContas

def clean_str(val):
    if val is None:
        return ""
    return str(val).replace("&nbsp;", " ").replace("\xa0", " ").strip()

def parse_decimal(val):
    if val is None or val == "":
        return Decimal("0.00")
    if isinstance(val, (int, float)):
        return Decimal(str(val)).quantize(Decimal("0.01"))
    if isinstance(val, Decimal):
        return val.quantize(Decimal("0.01"))
    s = clean_str(val).replace("%", "").strip()
    if not s:
        return Decimal("0.00")
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s).quantize(Decimal("0.01"))
    except Exception:
        return Decimal("0.00")

def find_excel_file(keyword):
    kw = keyword.lower()
    search_dirs = [
        ROOT_DIR / "scripts",
        ROOT_DIR / "backups" / "fabio",
        Path("/app/scripts"),
        Path("/app/backups/fabio"),
    ]
    
    print(f"Buscando arquivo com a palavra-chave '{kw}'...")
    for d in search_dirs:
        if d.exists() and d.is_dir():
            for f in d.iterdir():
                if f.is_file() and f.suffix.lower() == ".xlsx":
                    f_name = f.name.lower()
                    if kw in f_name and "ifood" not in f_name:
                        print(f"  ✅ Encontrado: {f}")
                        return f
    return None

def run(dry_run=True):
    print("======================================================================")
    print("ALINHAMENTO HISTÓRICO DAS PLANILHAS GOOGLE SHEETS (ORIGEM='WEB')")
    print(f"Modo: {'SIMULAÇÃO (dry-run)' if dry_run else 'APLICAR NO BANCO (COMMIT)'}")
    print("======================================================================")
    
    files_map = {
        35: ("umarizal", "Umarizal"),
        37: ("ananindeua", "Ananindeua"),
        39: ("marco", "Salão"),
        40: ("marco", "Delivery"),
    }
    
    # Category code mapping
    code_mapping = {
        "01.01.": "01.01.01", # Dinheiro
        "01.01": "01.01.01",
        "01.02.": "01.01.02", # Cartão crédito
        "01.02": "01.01.02",
        "01.03.": "01.01.03", # Cartão débito
        "01.03": "01.01.03",
        "01.04.": "01.01.04", # Pix/depósito
        "01.04": "01.01.04",
        "01.05.": "01.01.05", # Pix QRS
        "01.05": "01.01.05",
        "01.06.": "01.01.06", # Outras receitas
        "01.06": "01.01.06",
    }
    
    db = Session(engine)
    try:
        for emp_id, (file_name, sheet_name) in files_map.items():
            xlsx_path = find_excel_file(file_name)
            if not xlsx_path:
                print(f"⚠️ Arquivo {file_name} não encontrado nas pastas do projeto. Ignorando empresa {emp_id}.")
                continue
                
            print(f"\nProcessing Empresa ID {emp_id} usando {xlsx_path.name}...")
            wb = openpyxl.load_workbook(xlsx_path, data_only=True)
            
            # Target sheet 'DRE'
            sheet = None
            for sname in wb.sheetnames:
                if "dre" in sname.lower() and "2025" not in sname and "2024" not in sname:
                    sheet = wb[sname]
                    print(f"  Aba DRE selecionada: '{sname}'")
                    break
            if not sheet:
                sheet = wb.active
                print(f"  Aba padrão selecionada: '{sheet.title}'")
            
            # Map header columns from Row 1
            month_cols = {}
            for col in range(4, 25): # Cols D (4), E (5), F (6)...
                cell_val = sheet.cell(1, col).value
                if not cell_val:
                    continue
                if isinstance(cell_val, (datetime, date)):
                    month_cols[col] = (cell_val.year, cell_val.month)
                    print(f"  Mês encontrado na Col {col} (Row 1): {cell_val.year}-{cell_val.month:02d}")
                else:
                    val = clean_str(cell_val)
                    if "/" in val:
                        parts = val.split("/")
                        if len(parts) == 2:
                            try:
                                m = int(parts[0])
                                y = int(parts[1])
                                if y < 100: y += 2000
                                month_cols[col] = (y, m)
                                print(f"  Mês encontrado na Col {col} (Row 1): {y}-{m:02d}")
                            except ValueError:
                                pass
                                
            print(f"  Total de colunas de meses identificadas em Row 1: {len(month_cols)}")
            
            # Load plano_contas for this company
            pcs = db.exec(select(PlanoContas).where(PlanoContas.empresa_id == emp_id, PlanoContas.is_deleted == False)).all()
            pc_by_code = {p.codigo: p for p in pcs if p.codigo}
            
            # Read rows for revenue categories
            updates_needed = []
            for row in range(2, 40):
                raw_code = clean_str(sheet.cell(row, 1).value)
                if not raw_code:
                    continue
                    
                # Clean code
                mapped_code = None
                for prefix, target in code_mapping.items():
                    if raw_code.startswith(prefix) or raw_code == prefix:
                        mapped_code = target
                        break
                        
                if not mapped_code or mapped_code not in pc_by_code:
                    continue
                    
                pc_obj = pc_by_code[mapped_code]
                
                for col, (y, m) in month_cols.items():
                    # Processar apenas o período histórico (até Junho/2026)
                    if y == 2026 and m > 6:
                        continue
                        
                    val = parse_decimal(sheet.cell(row, col).value)
                    comp_date = date(y, m, 1)
                    end_day = 31 if m in [1, 3, 5, 7, 8, 10, 12] else (30 if m != 2 else (29 if y % 4 == 0 else 28))
                    venc_date = date(y, m, end_day)
                    
                    updates_needed.append({
                        "empresa_id": emp_id,
                        "plano_contas_id": pc_obj.id,
                        "codigo": mapped_code,
                        "nome": pc_obj.nome,
                        "data_competencia": comp_date,
                        "data_vencimento": venc_date,
                        "valor": val
                    })
                    
            print(f"  Encontradas {len(updates_needed)} entradas de faturamento da planilha.")
            
            # Reconcile with existing WEB launches in lancamentos
            now = datetime.utcnow()
            for item in updates_needed:
                existing = db.exec(
                    select(Lancamento)
                    .where(
                        Lancamento.empresa_id == emp_id,
                        Lancamento.plano_contas_id == item["plano_contas_id"],
                        Lancamento.origem == "WEB",
                        Lancamento.data_competencia == item["data_competencia"],
                        Lancamento.is_deleted == False
                    )
                ).all()
                
                target_val = item["valor"]
                
                if existing:
                    # Update first item to exact spreadsheet value, soft delete rest if duplicates
                    first = existing[0]
                    if abs((first.valor_pago or first.valor_previsto or Decimal("0.00")) - target_val) > Decimal("0.01"):
                        print(f"  [AJUSTE] {item['codigo']} em {item['data_competencia']}: R$ {first.valor_previsto} -> R$ {target_val}")
                        if not dry_run:
                            first.valor_previsto = target_val
                            first.valor_pago = target_val
                            first.updated_at = now
                            db.add(first)
                            
                    for dup in existing[1:]:
                        print(f"  [REMOVER DUPLICADO WEB] ID {dup.id} ({item['codigo']}): R$ {dup.valor_previsto}")
                        if not dry_run:
                            dup.is_deleted = True
                            dup.deleted_at = now
                            db.add(dup)
                else:
                    if target_val > Decimal("0.00"):
                        print(f"  [INSERIR FALTANTE WEB] {item['codigo']} em {item['data_competencia']}: R$ {target_val}")
                        if not dry_run:
                            new_lan = Lancamento(
                                empresa_id=emp_id,
                                plano_contas_id=item["plano_contas_id"],
                                tipo="R",
                                origem="WEB",
                                valor_previsto=target_val,
                                valor_pago=target_val,
                                data_competencia=item["data_competencia"],
                                data_vencimento=item["data_vencimento"],
                                data_pagamento=item["data_vencimento"],
                                status="PAGO",
                                observacao="Importação Financeiro Planilha Oficial",
                                is_deleted=False,
                                created_at=now,
                                updated_at=now
                            )
                            db.add(new_lan)
                            
        if not dry_run:
            db.commit()
            print("\n✅ Alinhamento concluído com sucesso!")
        else:
            print("\n🔍 Simulação de alinhamento finalizada. Nenhuma alteração realizada.")
            
    except Exception as e:
        print(f"Erro durante alinhamento: {e}")
        db.rollback()
    finally:
        db.close()

if __name__ == "__main__":
    dry_run = "--apply" not in sys.argv
    run(dry_run)

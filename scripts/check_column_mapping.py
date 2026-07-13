import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')
base = r'backups/fabio'

# What the script expects for each sheet
SCRIPT_EXPECTS = {
    "Tb_Movimentacao": ["IdPDV", "Data", "Tipo", "Historico", "FormaPagto", "Bandeira",
                        "Qtde Parcelas", "Valor Cheio", "Centro de Custo", "Usuario"],
    "Tb_Financeira":   ["IdFinanceiro", "Data Vcto", "Data Pagto", "Tipo", "Classificacao",
                        "Descricao", "Valor Previsto", "Valor Realizado", "Banco",
                        "Situacao", "IdParcelamento", "Interessado", "Centro de Custo"],
    "Tb_Ifood":        ["Id_Ifood", "Data", "Hora", "Forma Pagto",
                        "Valor Bruto", "Valor Liquido", "Status", "Data Recebimento"],
    "Tb_Cartoes":      ["Data Venda", "Forma Pagto", "Bandeira", "Parcela",
                        "Qtde Parcelas", "Valor", "% Taxa", "Juros"],
    "TxCartoes":       ["Tipo", "Parcelas", "Bandeira", "Taxa", "Dias"],
    "CadClassificacao": ["Classificacao", "Tipo"],
    "Tb_Banco":        ["Banco", "Saldo"],
}

def normalize(s):
    import unicodedata
    s = str(s).strip()
    s = unicodedata.normalize('NFD', s)
    s = ''.join(c for c in s if not unicodedata.combining(c))
    return s.lower()

FILES = {
    "Umarizal":    "Base_PizzaFabioUmarizal (1).xlsx",
    "Marco (Base)":"Base_PizzaFabioMarco.xlsx",
    "Marco iFood": "Base_IFood_PizzaFabioMarco (1).xlsx",
}

print("=" * 70)
print("COMPARACAO COLUNAS: Planilha vs Script")
print("=" * 70)

for file_label, filename in FILES.items():
    wb = openpyxl.load_workbook(base + '/' + filename, read_only=True, data_only=True)
    print("\n\n### " + file_label + " ###")
    for sheet_name, expected_cols in SCRIPT_EXPECTS.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        actual_headers = [str(c.value).strip() if c.value else '' for c in list(ws.rows)[0]]
        actual_norm = [normalize(h) for h in actual_headers]
        print("\n  [" + sheet_name + "]")
        print("  Colunas reais: " + str(actual_headers))
        print("  Verificacao:")
        for col in expected_cols:
            col_norm = normalize(col)
            found = col_norm in actual_norm
            if found:
                idx = actual_norm.index(col_norm)
                real_name = actual_headers[idx]
                match = " OK  " if real_name == col else " ~OK ('" + real_name + "')"
                print("    [" + match + "] " + col)
            else:
                # Find closest fallback (positional)
                pos = list(SCRIPT_EXPECTS[sheet_name]).index(col)
                fallback = actual_headers[pos] if pos < len(actual_headers) else "N/A"
                print("    [MISS] " + col + " => usaria fallback idx " + str(pos) + " = '" + fallback + "'")
    wb.close()

print("\n\nAnalisando relacionamentos PDV/iFood -> Financeiro...")
wb = openpyxl.load_workbook(base + '/Base_PizzaFabioMarco.xlsx', read_only=True, data_only=True)
for sheet in ["Tb_Financeira", "Tb_Ifood", "Tb_Cartoes"]:
    if sheet in wb.sheetnames:
        ws = wb[sheet]
        rows_sample = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            rows_sample.append([str(v)[:30] if v is not None else '' for v in row])
            if i >= 4: break
        print("\n  " + sheet + " (primeiras 4 linhas):")
        for r in rows_sample:
            print("    " + str(r))
wb.close()

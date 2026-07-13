import openpyxl, sys
from datetime import datetime
sys.stdout.reconfigure(encoding='utf-8')
base = r'backups/fabio'

def parse_date(val):
    if val is None:
        return None
    if isinstance(val, (datetime,)):
        return val
    try:
        from datetime import date
        if isinstance(val, date):
            return datetime(val.year, val.month, val.day)
    except Exception:
        pass
    try:
        return datetime.strptime(str(val).strip(), "%Y-%m-%d %H:%M:%S")
    except Exception:
        pass
    try:
        return datetime.strptime(str(val).strip()[:10], "%Y-%m-%d")
    except Exception:
        pass
    try:
        return datetime.strptime(str(val).strip(), "%d/%m/%Y")
    except Exception:
        pass
    return None

# ---- Umarizal: separar datas por centro de custo ----
print("=== Base_PizzaFabioUmarizal (1).xlsx - Tb_Movimentacao ===")
wb = openpyxl.load_workbook(base + '/Base_PizzaFabioUmarizal (1).xlsx', read_only=True, data_only=True)
ws = wb['Tb_Movimentacao']
rows = list(ws.iter_rows(values_only=True))
headers = [str(h).strip() if h else '' for h in rows[0]]
dt_idx = headers.index("Data") if "Data" in headers else 1
cc_idx = headers.index("Centro de Custo") if "Centro de Custo" in headers else 8

marco_dates = []
umarizal_dates = []
for row in rows[1:]:
    if not any(row): continue
    cc = str(row[cc_idx]).strip() if cc_idx < len(row) and row[cc_idx] else ''
    d = parse_date(row[dt_idx] if dt_idx < len(row) else None)
    if d:
        if 'marco' in cc.lower():
            marco_dates.append(d)
        elif 'umarizal' in cc.lower():
            umarizal_dates.append(d)
wb.close()

if marco_dates:
    print("Marco rows: " + str(len(marco_dates)) + " linhas | De: " + str(min(marco_dates).date()) + " Ate: " + str(max(marco_dates).date()))
if umarizal_dates:
    print("Umarizal rows: " + str(len(umarizal_dates)) + " linhas | De: " + str(min(umarizal_dates).date()) + " Ate: " + str(max(umarizal_dates).date()))

# ---- Marco Base: datas ----
print()
print("=== Base_PizzaFabioMarco.xlsx - Tb_Movimentacao ===")
wb2 = openpyxl.load_workbook(base + '/Base_PizzaFabioMarco.xlsx', read_only=True, data_only=True)
ws2 = wb2['Tb_Movimentacao']
rows2 = list(ws2.iter_rows(values_only=True))
headers2 = [str(h).strip() if h else '' for h in rows2[0]]
dt_idx2 = headers2.index("Data") if "Data" in headers2 else 1
marco2_dates = []
for row in rows2[1:]:
    if not any(row): continue
    d = parse_date(row[dt_idx2] if dt_idx2 < len(row) else None)
    if d:
        marco2_dates.append(d)
wb2.close()

if marco2_dates:
    print("Marco (Base) rows: " + str(len(marco2_dates)) + " linhas | De: " + str(min(marco2_dates).date()) + " Ate: " + str(max(marco2_dates).date()))

print()
print("=== CONCLUSAO ===")
if marco_dates and marco2_dates:
    last_marco_umarizal = max(marco_dates)
    first_marco_base = min(marco2_dates)
    last_marco_base = max(marco2_dates)
    print("Marco no Umarizal.xlsx termina em: " + str(last_marco_umarizal.date()))
    print("Marco no Marco.xlsx comeca em:     " + str(first_marco_base.date()))
    print("Marco no Marco.xlsx termina em:    " + str(last_marco_base.date()))
    if first_marco_base > last_marco_umarizal:
        print("=> Sem sobreposicao de datas: Marco.xlsx e CONTINUACAO apos separacao!")
    elif last_marco_base > last_marco_umarizal:
        print("=> Marco.xlsx tem dados MAIS RECENTES que o Umarizal.xlsx (separacao parcial)")
    else:
        print("=> Datas se sobrepoem: periodo compartilhado + separado")

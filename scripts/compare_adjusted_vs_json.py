# scripts/compare_adjusted_vs_json.py
import sys
import openpyxl
from pathlib import Path
from decimal import Decimal

ROOT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT_DIR))

def main():
    xlsx_path = ROOT_DIR / "scripts" / "Base_PizzaFabioUmarizal.xlsx"
    if not xlsx_path.exists():
        print(f"Arquivo {xlsx_path} não encontrado!")
        return
        
    print(f"Lendo planilha: {xlsx_path}")
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    
    # Print sheet names
    print("Abas da planilha:", wb.sheetnames)
    
    sheet = wb["Umarizal"] if "Umarizal" in wb.sheetnames else wb.active
    
    for row in range(1, 30):
        vals = [sheet.cell(row, col).value for col in range(1, 12)]
        if any(vals):
            print(f"Linha {row:2d}: {vals}")

if __name__ == "__main__":
    main()

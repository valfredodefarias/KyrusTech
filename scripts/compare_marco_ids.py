import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')
base = r'backups/fabio'

def get_ids(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        return set()
    ws = wb[sheet]
    ids = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            try:
                ids.add(int(float(str(row[0]))))
            except Exception:
                pass
    wb.close()
    return ids

for sheet in ['Tb_Financeira', 'Tb_Movimentacao', 'Tb_Ifood']:
    print('--- ' + sheet + ' ---')
    base_ids  = get_ids(base + '/Base_PizzaFabioMarco.xlsx', sheet)
    ifood_ids = get_ids(base + '/Base_IFood_PizzaFabioMarco (1).xlsx', sheet)
    overlap   = base_ids & ifood_ids
    mn_b = str(min(base_ids))  if base_ids  else '?'
    mx_b = str(max(base_ids))  if base_ids  else '?'
    mn_i = str(min(ifood_ids)) if ifood_ids else '?'
    mx_i = str(max(ifood_ids)) if ifood_ids else '?'
    print('Marco Base:  ' + str(len(base_ids))  + ' ids | min=' + mn_b + ' max=' + mx_b)
    print('Marco iFood: ' + str(len(ifood_ids)) + ' ids | min=' + mn_i + ' max=' + mx_i)
    print('Sobreposicao: ' + str(len(overlap)) + ' ids em COMUM')
    print('Exclusivos Base: ' + str(len(base_ids - ifood_ids)) + ' | Exclusivos iFood: ' + str(len(ifood_ids - base_ids)))
    print()

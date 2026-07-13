# scripts/verify_dre_match.py
import sys
import os
import openpyxl
from datetime import date, datetime
from decimal import Decimal
from sqlmodel import Session, select, func

sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app.db.session import Session, engine
from app.models.lancamento import Lancamento
from app.models.empresa import Empresa

def clean_str(val):
    if val is None:
        return ""
    return str(val).strip()

def parse_decimal(val):
    if val is None or val == "":
        return Decimal("0.00")
    if isinstance(val, (int, float)):
        return Decimal(str(val))
    if isinstance(val, Decimal):
        return val
    try:
        return Decimal(str(val).replace(",", ".").strip())
    except:
        return Decimal("0.00")

def parse_date(val):
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if not val:
        return None
    s = clean_str(val)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None

def main():
    db = Session(engine)
    empresa = db.exec(select(Empresa).where(Empresa.nome_fantasia == "Pizza Fábio Umarizal")).first()
    if not empresa:
        print("Empresa Pizza Fábio Umarizal não encontrada no banco.")
        return
    
    empresa_id = empresa.id
    print(f"Verificando Empresa: {empresa.nome_fantasia} (ID: {empresa_id})")

    # Load spreadsheet
    file_path = "scripts/Base_PizzaFabioUmarizal.xlsx"
    if not os.path.exists(file_path):
        print(f"Planilha não encontrada em: {file_path}")
        return

    wb = openpyxl.load_workbook(file_path, data_only=True)
    if "Tb_Financeira" not in wb.sheetnames:
        print("Aba Tb_Financeira não encontrada na planilha.")
        return

    sheet = wb["Tb_Financeira"]
    rows = list(sheet.iter_rows(values_only=True))
    headers = [clean_str(h) for h in rows[0]]

    pag_idx = headers.index("Data Pagto")
    tipo_idx = headers.index("Tipo")
    real_idx = headers.index("Valor Realizado")
    sit_idx = headers.index("Situação")
    cc_idx = headers.index("Centro de Custo") if "Centro de Custo" in headers else 12

    # July 2024 spreadsheet sums
    sheet_receitas = Decimal("0.00")
    sheet_despesas = Decimal("0.00")

    print(f"Total de linhas na planilha: {len(rows)}")
    print("Primeiras 5 datas:")
    for row in rows[1:6]:
        print(f"  Data Pagto: {row[pag_idx]} | Centro de Custo: {row[cc_idx]} | Situação: {row[sit_idx]} (tipo: {type(row[sit_idx])})")

    matching_count = 0
    for row in rows[1:]:
        if not any(row):
            continue
        row_cc = clean_str(row[cc_idx]) if cc_idx < len(row) else ""
        if row_cc.lower() != "umarizal":
            continue
        sit = clean_str(row[sit_idx])
        if sit.lower() != "pago":
            continue
            
        dt_val = row[pag_idx]
        dt = parse_date(dt_val)
        if not dt:
            continue
            
        matching_count += 1
        if matching_count <= 5:
            print(f"Row {matching_count}: date={dt}, val={row[real_idx]}, tipo={row[tipo_idx]}, class={row[headers.index('Classificação')]}")

        if dt.year == 2024 and dt.month == 7:
            val = parse_decimal(row[real_idx])
            tipo = clean_str(row[tipo_idx]).lower()
            
            if tipo == "recebimento":
                sheet_receitas += abs(val)
            else:
                sheet_despesas += abs(val)

    wb.close()

    # July 2024 database sums (excluding PDV origin to match only Tb_Financeira first)
    db_receitas = db.exec(
        select(func.sum(Lancamento.valor_pago))
        .where(
            Lancamento.empresa_id == empresa_id,
            Lancamento.tipo == "RECEITA",
            Lancamento.status == "PAGO",
            Lancamento.origem == "WEB",
            Lancamento.data_pagamento >= date(2024, 7, 1),
            Lancamento.data_pagamento <= date(2024, 7, 31),
            Lancamento.is_deleted == False
        )
    ).first() or Decimal("0.00")

    db_despesas = db.exec(
        select(func.sum(Lancamento.valor_pago))
        .where(
            Lancamento.empresa_id == empresa_id,
            Lancamento.tipo == "DESPESA",
            Lancamento.status == "PAGO",
            Lancamento.origem == "WEB",
            Lancamento.data_pagamento >= date(2024, 7, 1),
            Lancamento.data_pagamento <= date(2024, 7, 31),
            Lancamento.is_deleted == False
        )
    ).first() or Decimal("0.00")

    print("\n--- COMPARAÇÃO DRE: JULHO DE 2024 ---")
    print(f"Receitas Planilha: R$ {sheet_receitas:,.2f} | Kyrus Banco: R$ {db_receitas:,.2f} | Dif: R$ {sheet_receitas - db_receitas:,.2f}")
    print(f"Despesas Planilha:  R$ {sheet_despesas:,.2f} | Kyrus Banco: R$ {db_despesas:,.2f} | Dif: R$ {sheet_despesas - db_despesas:,.2f}")
    
    # Net profit comparison
    sheet_net = sheet_receitas - sheet_despesas
    db_net = db_receitas - db_despesas
    print(f"Resultado Líquido Planilha: R$ {sheet_net:,.2f} | Kyrus Banco: R$ {db_net:,.2f} | Dif: R$ {sheet_net - db_net:,.2f}")

    if abs(sheet_net - db_net) < Decimal("0.01"):
        print("\n✅ SUCESSO! A DRE do Kyrus ERP bate 100% com os dados consolidados da planilha original!")
    else:
        print("\n❌ AVISO: Há discrepâncias entre os totais.")

if __name__ == "__main__":
    main()

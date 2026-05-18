from openpyxl import load_workbook
p = r"c:\Users\Ciro\Documents\ERP\KyrusERP\xlsx\LinkFinanceiro\classificacao.xlsx"
wb = load_workbook(p, data_only=True)
print("SHEETS", wb.sheetnames)
for sh in wb.sheetnames:
    ws = wb[sh]
    print("---", sh, "rows", ws.max_row, "cols", ws.max_column)
    max_row = min(15, ws.max_row)
    for r in ws.iter_rows(min_row=1, max_row=max_row, values_only=True):
        print(r)
